"""Tests with made-up data only. Every folder is temporary; the database is never touched."""
import asyncio
import os
import re
import shutil
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

TEMP = Path(tempfile.mkdtemp(prefix="resume-builder-test-"))
os.environ["RESUME_OUTPUT_DIR"] = str(TEMP / "user-builds")
os.environ["RESUME_DATA_DIR"] = str(TEMP / "data")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import builds  # noqa: E402
import config  # noqa: E402
import documents  # noqa: E402
import jobfinder_db  # noqa: E402
import profile_import  # noqa: E402
import references  # noqa: E402
import resume_file  # noqa: E402
import writing_rules  # noqa: E402
from models import ResumeContent, merge_layout  # noqa: E402
from pdf_render import render_resume  # noqa: E402
from contextlib import contextmanager  # noqa: E402


@contextmanager
def _no_real_database(commit=False):
    """Tests must never reach the user's real Job Finder database. Tests that need one patch _cursor."""
    raise jobfinder_db.JobFinderUnavailable("tests have no database")
    yield  # pragma: no cover


jobfinder_db._cursor = _no_real_database

PROFILE = {"first_name": "Jane", "last_name": "Doe", "home_location": "Springfield, MD", "state": "MD",
           "primary_job_title": "Web Developer", "job_titles": ["Web Developer"], "skills": ["HTML", "CSS"],
           "work_history": [{"company": "Example Co", "role": "Web Developer", "dates": "2020 - Present", "description": ""}]}
JOB = {"job_id": 7, "job_title": "Front-End Web Developer", "company": "Acme Widgets", "location": "Springfield, MD",
       "work_arrangement": "Hybrid", "application_status": "None", "saved": True, "date_found": "2026-09-30",
       "description": "Build pages with HTML and CSS.", "salary": "", "posted": "", "listing_location": "",
       "listing_skills": ["HTML", "CSS"], "url": "https://example.com/job/7", "notes": ""}
SAMPLE = {
    "full_name": "Jane Doe", "headline": "Web Developer",
    "contact": ["jane@example.com", "555-0100", "Springfield, MD"],
    "sections": [
        {"title": "Summary", "text": "Web developer with experience building accessible sites."},
        {"title": "Experience", "items": [{"heading": "Web Developer", "subheading": "Example Co",
                                           "location": "Springfield, MD", "dates": "2020 – Present",
                                           "bullets": ["Built pages in HTML & CSS", "Improved <accessibility> scores"]}]},
        {"title": "Skills", "text": "HTML · CSS · JavaScript"},
    ],
}
LETTER = {"full_name": "Jane Doe", "contact": ["jane@example.com"], "recipient": ["Hiring Manager", "Acme Widgets"],
          "paragraphs": ["I am applying for the role.", "Thank you for your time."]}
SAMPLE_LAYOUT = {"font_family": "Georgia", "font_kind": "serif", "body_size": 10.5, "name_size": 24, "heading_size": 12.5,
                 "accent_color": "#2E5A88", "text_color": "#222222", "name_align": "center",
                 "heading_case": "upper", "heading_rule": True, "margin_in": 0.7}


def make_sample_pdf():
    path = TEMP / "jane-sample.pdf"
    render_resume(ResumeContent.model_validate(SAMPLE), merge_layout(SAMPLE_LAYOUT), path)
    return path.read_bytes()


def make_sample_docx():
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run("Jane Doe")
    run.font.size, run.bold, run.font.color.rgb = Pt(22), True, RGBColor(0x1F, 0x6F, 0x4F)
    document.add_paragraph("jane@example.com | 555-0100")
    for title, body in (("EXPERIENCE", "Web Developer, Example Co, 2020 - Present"), ("SKILLS", "HTML, CSS")):
        heading = document.add_paragraph()
        heading_run = heading.add_run(title)
        heading_run.bold, heading_run.font.size = True, Pt(13)
        heading_run.font.color.rgb = RGBColor(0x1F, 0x6F, 0x4F)
        document.add_paragraph(body)
    path = TEMP / "jane-sample.docx"
    document.save(path)
    return path.read_bytes()


class ResumeBuilderTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()

    def test_paths_are_temporary(self):
        self.assertTrue(str(config.OUTPUT_DIR).startswith(str(TEMP)))
        self.assertTrue(str(config.DATA_DIR).startswith(str(TEMP)))

    def test_pdf_upload_measures_layout(self):
        info = resume_file.save_upload("Jane Resume.pdf", make_sample_pdf())
        layout = info["layout"]
        self.assertEqual(layout["font_family"], "Georgia")
        self.assertEqual(layout["font_kind"], "serif")
        self.assertEqual(layout["name_align"], "center")
        self.assertTrue(layout["heading_rule"])
        self.assertEqual(layout["heading_case"], "upper")
        self.assertEqual(layout["accent_color"], "#2E5A88")
        self.assertAlmostEqual(layout["body_size"], 10.5, delta=0.5)
        self.assertAlmostEqual(layout["name_size"], 24, delta=0.5)
        self.assertIn("Example Co", resume_file.current_text())
        self.assertEqual(len(resume_file.current_page_images()), 1)

    def test_docx_upload_without_word_uses_styles(self):
        with mock.patch.object(resume_file, "docx_to_pdf", side_effect=OSError("no Word")):
            info = resume_file.save_upload("jane.docx", make_sample_docx())
        layout = info["layout"]
        self.assertEqual(layout["font_family"], "Calibri")
        self.assertEqual(layout["name_align"], "center")
        self.assertEqual(layout["accent_color"], "#1F6F4F")
        self.assertEqual(layout["heading_case"], "upper")
        self.assertTrue(info["notes"])
        self.assertIn("Example Co", resume_file.current_text())

    @unittest.skipUnless(os.environ.get("RESUME_TEST_WORD") == "1", "set RESUME_TEST_WORD=1 to convert with Word")
    def test_docx_upload_with_word(self):
        info = resume_file.save_upload("jane.docx", make_sample_docx())
        self.assertEqual(info["notes"], [])
        self.assertEqual(info["layout"]["font_family"], "Calibri")
        self.assertEqual(len(resume_file.current_page_images()), 1)

    def test_rejects_wrong_files(self):
        for name, data in (("a.txt", b"hi"), ("a.pdf", b"PK not a pdf"), ("a.docx", b"%PDF-1.4"), ("a.doc", b"x")):
            with self.assertRaises(resume_file.ResumeFileError):
                resume_file.save_upload(name, data)
        self.assertIsNone(resume_file.current_info())

    def test_file_names(self):
        today = date.today().strftime("%m-%d-%Y")
        first = builds.save_build("resume", SAMPLE, job_title="Front-End Web Developer", profile=PROFILE)
        second = builds.save_build("resume", SAMPLE, job_title="Front-End Web Developer", profile=PROFILE)
        letter = builds.save_build("cover_letter", LETTER, job_title="Front-End Web Developer", profile=PROFILE)
        self.assertEqual(first["pdf"], f"Jane_Doe_Front_End_Web_Developer_{today}.pdf")
        self.assertEqual(second["pdf"], f"Jane_Doe_Front_End_Web_Developer_{today}_2.pdf")
        self.assertEqual(letter["pdf"], f"Jane_Doe_Front_End_Web_Developer_Cover_Letter_{today}.pdf")
        for record in (first, second, letter):
            self.assertTrue((config.OUTPUT_DIR / record["pdf"]).read_bytes().startswith(b"%PDF"))

    def test_file_names_carry_the_company_and_never_say_ats(self):
        when = date(2026, 9, 23)
        self.assertEqual(builds.build_filename("Jane", "Doe", "resume", "Web Designer", "Pacvue", when),
                         "Jane_Doe_Pacvue_Web_Designer_09-23-2026.pdf")
        self.assertEqual(builds.build_filename("Jane", "Doe", "resume", "ATS Web Designer", "Pacvue", when),
                         "Jane_Doe_Pacvue_Web_Designer_09-23-2026.pdf")
        self.assertEqual(builds.build_filename("Jane", "Doe", "cover_letter", "Web Designer", "Pacvue", when),
                         "Jane_Doe_Pacvue_Web_Designer_Cover_Letter_09-23-2026.pdf")

    def test_name_falls_back_to_resume_name(self):
        record = builds.save_build("resume", SAMPLE, job_title="Analyst", profile={})
        self.assertTrue(record["pdf"].startswith("Jane_Doe_Analyst_"))

    def test_uses_uploaded_layout_then_overrides(self):
        resume_file.save_upload("Jane Resume.pdf", make_sample_pdf())
        record = builds.save_build("resume", SAMPLE, job_title="Dev", profile=PROFILE, layout={"accent_color": "#AA0000"})
        self.assertEqual(record["layout"]["font_family"], "Georgia")
        self.assertEqual(record["layout"]["accent_color"], "#AA0000")

    def test_replace_keeps_file_name(self):
        record = builds.save_build("resume", SAMPLE, job_title="Dev", profile=PROFILE)
        changed = dict(SAMPLE, headline="Senior Web Developer")
        again = builds.save_build("resume", changed, job_title="Dev", profile=PROFILE, replace=record["pdf"])
        self.assertEqual(again["pdf"], record["pdf"])
        self.assertEqual(builds.load_draft(record["pdf"])["content"]["headline"], "Senior Web Developer")
        self.assertEqual(len(builds.list_builds()), 1)

    def test_references_print_last_in_resume(self):
        refs = [{"name": "Sam Lee", "relationship": "Supervisor", "user_job": "Web Developer at Example Co",
                 "job_title": "IT Director", "company": "Example Co", "phone": "555-0111", "email": "sam@example.com"}]
        record = builds.save_build("resume", dict(SAMPLE, references=refs), job_title="Dev", profile=PROFILE)
        text = resume_file.pdf_text(config.OUTPUT_DIR / record["pdf"])
        self.assertIn("REFERENCES", text)
        self.assertGreater(text.index("REFERENCES"), text.index("SKILLS"))
        self.assertIn("Supervisor while I was Web Developer at Example Co", text)
        self.assertIn("555-0111", text)
        plain = builds.save_build("resume", SAMPLE, job_title="Dev", profile=PROFILE)
        self.assertNotIn("REFERENCES", resume_file.pdf_text(config.OUTPUT_DIR / plain["pdf"]))

    def test_rejects_unsafe_names(self):
        for bad in ("../secret.pdf", "..\\x.pdf", "a/b.pdf", "x.json"):
            with self.assertRaises(ValueError):
                builds.load_draft(bad)


class WebAppTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()
        import app as web
        self.web = web
        web.app.config["TESTING"] = True
        self.client = web.app.test_client()
        patcher = mock.patch.object(jobfinder_db, "get_profile", return_value=PROFILE)
        patcher.start()
        self.addCleanup(patcher.stop)

    def csrf(self):
        page = self.client.get("/").get_data(as_text=True)
        return re.search(r'name="csrf" value="([^"]+)"', page).group(1)

    def test_index_and_upload(self):
        self.assertIn("Jane Doe", self.client.get("/").get_data(as_text=True))
        from io import BytesIO
        response = self.client.post("/upload", data={"csrf": self.csrf(), "resume": (BytesIO(make_sample_pdf()), "jane.pdf")},
                                    content_type="multipart/form-data", follow_redirects=True)
        page = response.get_data(as_text=True)
        self.assertIn("Uploaded jane.pdf.", page)
        self.assertNotIn("Measured layout", page)  # still measured and used, just not shown
        self.assertEqual(resume_file.current_info()["layout"]["font_family"], "Georgia")
        self.assertEqual(self.client.get("/resume/page/0.png").mimetype, "image/png")

    def test_profile_form_saves_everything(self):
        with mock.patch.object(jobfinder_db, "get_profile", return_value=dict(PROFILE, first_name="", last_name="")):
            page = self.client.get("/").get_data(as_text=True)
        self.assertIn("Add your name and location", page)
        self.assertIn('<details class="name-edit" open>', page)
        form = {"csrf": self.csrf(), "first_name": " Jane ", "last_name": "Doe", "home_location": "Springfield, MD",
                "primary_job_title": "Web Developer", "other_titles": "Front End Developer,\nUX Designer",
                "skills": "html, CSS\nFigma", "wh_count": "2",
                "wh0_role": "Web Developer", "wh0_company": "Example Co", "wh0_dates": "2020 - Present", "wh0_description": "Built pages",
                "wh0_supervisor_name": "Sam Lee",
                "wh1_role": "", "wh1_company": ""}
        with mock.patch.object(jobfinder_db, "save_profile") as save:
            page = self.client.post("/profile", data=form, follow_redirects=True).get_data(as_text=True)
        save.assert_called_once()
        args = save.call_args.args
        self.assertEqual(args[:5], ("Jane", "Doe", "Springfield, MD", "Web Developer", ["Front End Developer", "UX Designer"]))
        self.assertEqual(args[5], ["html", "CSS", "Figma"])
        self.assertEqual({k: v for k, v in args[6][0].items() if v}, {"role": "Web Developer", "company": "Example Co", "dates": "2020 - Present",
                                                                       "description": "Built pages", "supervisor_name": "Sam Lee"})
        self.assertIn("Saved to your Job Finder profile.", page)
        with mock.patch.object(jobfinder_db, "save_profile") as save:
            page = self.client.post("/profile", data={"csrf": self.csrf(), "first_name": "Jane", "last_name": ""},
                                    follow_redirects=True).get_data(as_text=True)
        save.assert_not_called()
        self.assertIn("Enter both a first and last name.", page)

    def test_save_profile_writes_like_job_finder(self):
        calls = []

        class Cursor:
            def execute(self, sql, params=()):
                calls.append((" ".join(sql.split()), params))

            def fetchone(self):
                return {"table": "user_profile"}

            def fetchall(self):  # SHOW COLUMNS: this database has only some of the optional detail columns yet
                return [{"Field": f} for f in ("id", "profile_id", "company", "role", "dates", "description", "street", "supervisor_name")]

        from contextlib import contextmanager

        @contextmanager
        def fake_cursor(commit=False):
            self.assertTrue(commit)
            yield Cursor()

        history = [{"role": "Web Developer", "company": "Example Co", "dates": "2020", "description": "x",
                    "street": "1 Main St", "supervisor_name": "Sam Lee", "phone": "555-0100"},
                   {"role": "", "company": "", "dates": "2019", "description": "dropped"}]
        with mock.patch.object(jobfinder_db, "_cursor", fake_cursor):
            jobfinder_db.save_profile(" Jane ", "Doe", "Bel Air, MD ", "Web Designer", ["web designer", "Web Producer"],
                                      ["html", "HTML", "Figma"], history)
        profile = [c for c in calls if c[0].startswith("UPDATE user_profile")][0]
        self.assertEqual(profile[1], ("Jane", "Doe", "Bel Air, MD", "MD", "Web Designer"))
        titles = [c[1][0] for c in calls if "user_profile_job_titles (profile_id" in c[0]]
        self.assertEqual(titles, ["Web Designer", "Web Producer"])
        skills = [c[1][0] for c in calls if "user_profile_skills (profile_id" in c[0]]
        self.assertEqual(skills, ["HTML", "Figma"])  # Job Finder's own tidying
        jobs = [c[1] for c in calls if "user_profile_work_history (profile_id" in c[0]]
        self.assertEqual(jobs, [["Example Co", "Web Developer", "2020", "x", "1 Main St", "Sam Lee"]])  # phone: no column yet
        self.assertFalse(any("user_profile_cities" in c[0] or "avatar" in c[0] for c in calls))

    def test_job_title_help(self):
        matches = self.client.get("/job-title-matches?q=web%20des").get_json()["matches"]
        self.assertEqual(matches[0], "Web Designer")
        suggested = self.client.get("/job-title-suggestions?titles=Web%20Designer").get_json()["suggestions"]
        self.assertTrue(suggested)
        self.assertNotIn("Web Designer", suggested)
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn('data-title-suggest="list"', page)
        self.assertIn("titles.js", page)

    def test_post_without_csrf_is_refused(self):
        self.assertEqual(self.client.post("/upload").status_code, 400)

    def test_edit_round_trip(self):
        references.add({"name": "Pat Kim", "relationship": "CEO / Owner", "notes": "private note"})
        record = builds.save_build("resume", SAMPLE, job_title="Dev", profile=PROFILE)
        url = f"/build/{record['pdf']}"
        page = self.client.get(url).get_data(as_text=True)
        self.assertIn("Example Co", page)
        self.assertIn("Pat Kim (CEO / Owner)", page)  # saved reference offered in the picker
        self.assertNotIn("private note", page)
        form = {"csrf": self.csrf(), "full_name": "Jane Doe", "headline": "Edited", "contact": "jane@example.com\n555-0100",
                "section_count": "2", "s0_title": "Experience", "s0_text": "", "s0_count": "2",
                "s0_i0_heading": "Web Developer", "s0_i0_subheading": "Example Co", "s0_i0_dates": "2020",
                "s0_i0_bullets": "One\n\nTwo", "s0_i1_heading": "", "s1_title": "", "s1_count": "1",
                "font_family": "Arial", "body_size": "11", "accent_color": "#123456", "heading_rule": "on",
                "ref_count": "2", "ref0_name": "Sam Lee", "ref0_relationship": "Co-worker", "ref1_name": ""}
        self.client.post(url, data=form)
        draft = builds.load_draft(record["pdf"])
        self.assertEqual(draft["content"]["headline"], "Edited")
        self.assertEqual(draft["content"]["sections"][0]["items"][0]["bullets"], ["One", "Two"])
        self.assertEqual(len(draft["content"]["sections"]), 1)
        self.assertEqual([r["name"] for r in draft["content"]["references"]], ["Sam Lee"])
        self.assertEqual(draft["layout"]["accent_color"], "#123456")
        self.assertEqual(draft["layout"]["font_family"], "Arial")
        with self.client.get(f"/files/{record['pdf']}") as response:
            self.assertEqual(response.mimetype, "application/pdf")

    def test_letter_edit_and_delete(self):
        record = builds.save_build("cover_letter", LETTER, job_title="Dev", profile=PROFILE)
        url = f"/build/{record['pdf']}"
        token = self.csrf()
        self.client.post(url, data={"csrf": token, "full_name": "Jane Doe", "body": "First.\r\n\r\nSecond.\r\nStill second.",
                                    "greeting": "Hello,", "closing": "Best,", "font_family": "Calibri"})
        self.assertEqual(builds.load_draft(record["pdf"])["content"]["paragraphs"], ["First.", "Second.\nStill second."])
        self.client.post(f"{url}/delete", data={"csrf": token})
        self.assertEqual(builds.list_builds(), [])
        self.assertFalse((config.OUTPUT_DIR / record["pdf"]).exists())

    def test_references_add_edit_delete(self):
        page = self.client.get("/").get_data(as_text=True)
        self.assertIn('value="Web Developer at Example Co"', page)  # suggestion from work history
        self.assertIn("No references yet.", page)
        token = self.csrf()
        sam = {"csrf": token, "name": "Sam Lee", "relationship": "Supervisor", "user_job": "Web Developer at Example Co",
               "job_title": "IT Director", "company": "Example Co", "phone": "555-0111", "email": "sam@example.com",
               "notes": "prefers a text first"}
        page = self.client.post("/references/add", data=sam, follow_redirects=True).get_data(as_text=True)
        self.assertIn("Added Sam Lee.", page)
        self.assertIn("Your job together: Web Developer at Example Co", page)
        self.assertIn('href="tel:555-0111"', page)
        page = self.client.post("/references/add", data={"csrf": token, "relationship": "Co-worker"},
                                follow_redirects=True).get_data(as_text=True)
        self.assertIn("Not added", page)
        self.client.post("/references/add", data={"csrf": token, "name": "Pat Kim", "email": "pat@example.com"})
        saved = references.load()
        self.assertEqual([r["name"] for r in saved], ["Sam Lee", "Pat Kim"])
        self.client.post(f"/references/{saved[0]['id']}/update", data=dict(sam, phone="555-0199"))
        self.assertEqual(references.load()[0]["phone"], "555-0199")
        self.client.post(f"/references/{saved[1]['id']}/delete", data={"csrf": token})
        self.assertEqual([r["name"] for r in references.load()], ["Sam Lee"])
        self.assertEqual(self.client.post("/references/nope/delete", data={"csrf": token}).status_code, 404)

    def test_files_outside_output_are_refused(self):
        self.assertEqual(self.client.get("/files/..%5Capp.py").status_code, 404)
        self.assertEqual(self.client.get("/files/../app.py").status_code, 404)
        self.assertEqual(self.client.get("/build/..%5Cx.pdf").status_code, 404)


class ConnectorTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()
        for name, value in (("get_profile", PROFILE), ("get_job", JOB), ("list_jobs", [JOB])):
            patcher = mock.patch.object(jobfinder_db, name, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def call(self, name, args=None):
        from mcp import Client
        import mcp_server

        async def run():
            async with Client(mcp_server.server, raise_exceptions=True) as client:
                if name == "__list__":
                    return await client.list_tools()
                return await client.call_tool(name, args or {})
        return asyncio.run(run())

    def text(self, result):
        return "\n".join(block.text for block in result.content if block.type == "text")

    def test_tools_listed(self):
        names = {tool.name for tool in self.call("__list__").tools}
        self.assertEqual(names, {"get_job_finder_profile", "get_current_resume", "get_references",
                                 "list_documents", "get_document", "list_saved_jobs", "list_apply_queue", "get_job",
                                 "save_resume", "save_cover_letter", "list_builds", "get_build", "get_writing_rules"})

    def test_an_older_single_rules_file_is_split_into_the_two_lists(self):
        config.DATA_DIR.mkdir(parents=True, exist_ok=True)
        (config.DATA_DIR / "writing-rules.md").write_text("Resume rule.\n\n## Cover Letter Guidelines\n- Short.\n", encoding="utf-8")
        self.assertEqual(writing_rules.load("resume"), "Resume rule.\n")
        self.assertEqual(writing_rules.load("cover_letter"), "## Cover Letter Guidelines\n- Short.\n")
        self.assertFalse((config.DATA_DIR / "writing-rules.md").exists())

    def test_writing_rules_are_read_first_and_saved_only_here(self):
        self.assertIn("No writing rules saved", self.text(self.call("get_writing_rules")))
        writing_rules.save("resume", "Use Technical Skills instead of a Professional Summary.\r\n")
        writing_rules.save("cover_letter", "Keep it short.")
        out = self.text(self.call("get_writing_rules"))
        self.assertIn("# Résumé Rules\n\nUse Technical Skills instead of a Professional Summary.", out)
        self.assertIn("# Cover Letter Rules\n\nKeep it short.", out)
        self.assertTrue((config.DATA_DIR / "resume-rules.md").exists())
        self.assertTrue((config.DATA_DIR / "cover-letter-rules.md").exists())
        import mcp_server
        self.assertIn("get_writing_rules FIRST", mcp_server.INSTRUCTIONS)

    def test_apply_queue_lists_what_each_job_still_needs(self):
        queue = [{"job_id": 7, "job_title": "Senior Product Designer", "company": "Owner"},
                 {"job_id": 9, "job_title": "Web Designer", "company": "Perplexity"}]
        saved = [{"job_id": 9, "kind": "resume", "pdf_exists": True}, {"job_id": 9, "kind": "cover_letter", "pdf_exists": False}]
        with mock.patch.object(jobfinder_db, "list_apply_queue", return_value=queue), \
                mock.patch.object(builds, "list_builds", return_value=saved):
            out = self.text(self.call("list_apply_queue"))
        self.assertRegex(out, r'"job_id": 7,[\s\S]*"needs": \[\s*"resume",\s*"cover_letter"\s*\]')
        self.assertRegex(out, r'"job_id": 9,[\s\S]*"needs": \[\s*"cover_letter"\s*\]')

    def test_current_resume_includes_page_picture(self):
        self.assertIn("No resume uploaded", self.text(self.call("get_current_resume")))
        resume_file.save_upload("jane.pdf", make_sample_pdf())
        result = self.call("get_current_resume")
        self.assertIn("Georgia", self.text(result))
        self.assertEqual([b.type for b in result.content].count("image"), 1)

    def test_save_resume_and_letter_by_job_id(self):
        today = date.today().strftime("%m-%d-%Y")
        out = self.text(self.call("save_resume", {"content": SAMPLE, "job_id": 7}))
        self.assertIn(f"Jane_Doe_Acme_Widgets_Front_End_Web_Developer_{today}.pdf", out)
        out = self.text(self.call("save_cover_letter", {"content": LETTER, "job_id": 7}))
        self.assertIn(f"Jane_Doe_Acme_Widgets_Front_End_Web_Developer_Cover_Letter_{today}.pdf", out)
        listed = self.text(self.call("list_builds"))
        self.assertIn("Acme Widgets", listed)

    def test_bad_layout_is_refused(self):
        result = self.call("save_resume", {"content": SAMPLE, "job_title": "Dev", "layout": {"accent_color": "red"}})
        self.assertTrue(result.is_error)
        self.assertEqual(builds.list_builds(), [])

    def test_references_without_private_notes(self):
        self.assertIn("No references saved", self.text(self.call("get_references")))
        references.add({"name": "Sam Lee", "relationship": "Supervisor", "user_job": "Web Developer at Example Co",
                        "phone": "555-0111", "notes": "secret note"})
        out = self.text(self.call("get_references"))
        self.assertIn("Sam Lee", out)
        self.assertNotIn('"id"', out)
        self.assertIn("Web Developer at Example Co", out)
        self.assertNotIn("secret note", out)

    def test_database_down_message(self):
        with mock.patch.object(jobfinder_db, "get_profile", side_effect=jobfinder_db.JobFinderUnavailable("DB down")):
            self.assertIn("DB down", self.text(self.call("get_job_finder_profile")))


SAMPLE_RESUME_TEXT = """Jane Doe​
(555) 010-0100 • jane@example.com
Springfield, MD
Professional Summary
Web developer who builds accessible sites.
Employment History
Web Designer​
Example Co​
Remote​
January 2025 - May 2025​
●
Rebuilt pages with Bootstrap, improving
consistency.
●
Learned the house style in two weeks.
Front End Web Developer
Sample Corp
Baltimore, MD
May 2015 – October 2024
●
Built landing pages in HTML and CSS.
Education History
State University
Towson, MD
Bachelor of Science, Graphic Design
"""


class ProfileImportTests(unittest.TestCase):
    def test_extract(self):
        found = profile_import.extract(SAMPLE_RESUME_TEXT)
        self.assertEqual(found["name"], "Jane Doe")
        self.assertEqual(found["location"], "Springfield, MD")
        self.assertIn("HTML", found["skills"])
        self.assertEqual([(j["role"], j["company"], j["dates"]) for j in found["work_history"]],
                         [("Web Designer", "Example Co", "2025-01 – 2025-05"),  # the month style the Dashboard's pickers use
                          ("Front End Web Developer", "Sample Corp", "2015-05 – 2024-10")])
        self.assertEqual(found["work_history"][0]["description"],
                         "• Rebuilt pages with Bootstrap, improving consistency.\n• Learned the house style in two weeks.")
        self.assertEqual(found["work_history"][1]["description"], "• Built landing pages in HTML and CSS.")

    def test_new_jobs_are_the_resumes_jobs_the_profile_lacks(self):
        empty = {"work_history": []}
        self.assertEqual([j["role"] for j in profile_import.new_jobs(SAMPLE_RESUME_TEXT, empty)],
                         ["Web Designer", "Front End Web Developer"])
        one_known = {"work_history": [{"role": "web designer", "company": "EXAMPLE CO", "dates": "", "description": ""}]}
        self.assertEqual([j["role"] for j in profile_import.new_jobs(SAMPLE_RESUME_TEXT, one_known)], ["Front End Web Developer"])
        everything = {"work_history": [{"role": j["role"], "company": j["company"], "dates": "", "description": ""}
                                       for j in profile_import.new_jobs(SAMPLE_RESUME_TEXT, empty)]}
        self.assertEqual(profile_import.new_jobs(SAMPLE_RESUME_TEXT, everything), [])

    def test_the_fill_button_adds_the_resumes_jobs_only_when_clicked(self):
        import app as web
        profile = {"first_name": "Jane", "last_name": "Doe", "home_location": "Springfield, MD", "state": "",
                   "primary_job_title": "", "job_titles": [], "skills": [], "work_history": []}
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "token"
        with mock.patch.object(jobfinder_db, "get_profile", return_value=profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save, \
                mock.patch.object(resume_file, "current_text", return_value=SAMPLE_RESUME_TEXT):
            refused = client.post("/profile/fill-from-resume", data={})
            self.assertEqual(refused.status_code, 400)  # no form token: nothing is saved
            self.assertFalse(save.called)
            answered = client.post("/profile/fill-from-resume", data={"csrf": "token"})
        self.assertEqual(answered.status_code, 302)
        self.assertTrue(save.called)
        saved_history = save.call_args.args[-1]
        self.assertEqual([j["role"] for j in saved_history], ["Web Designer", "Front End Web Developer"])

    def test_the_header_shows_only_the_version_under_the_title(self):
        html = (Path(__file__).resolve().parents[1] / "templates" / "index.html").read_text(encoding="utf-8")
        self.assertIn('<h1>Résumé Builder</h1>\n  <p class="app-version">Version {{ app_version }}</p>', html)
        self.assertNotIn("Claude Desktop writes a", html)

    def test_the_profile_page_offers_the_resumes_jobs_with_a_button(self):
        html = (Path(__file__).resolve().parents[1] / "templates" / "index.html").read_text(encoding="utf-8")
        self.assertIn("{% if resume_jobs %}", html)
        self.assertIn("Add them to my Profile", html)
        self.assertIn("fill_profile_from_resume", html)

    def fill(self, profile):
        with mock.patch.object(jobfinder_db, "get_profile", return_value=profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            added = profile_import.fill_profile(SAMPLE_RESUME_TEXT)
        return added, save

    def test_fills_an_empty_profile(self):
        empty = {"first_name": "", "last_name": "", "home_location": "", "state": "", "primary_job_title": "",
                 "job_titles": ["web designer", "Visual Designer"], "skills": [], "work_history": []}
        added, save = self.fill(empty)
        self.assertEqual(added[:2], ["name", "location"])
        self.assertIn("2 jobs", added)
        self.assertIn("primary job title", added)
        first, last, location, primary, others, skills, history = save.call_args.args
        self.assertEqual((first, last, location, primary, others), ("Jane", "Doe", "Springfield, MD", "web designer", ["Visual Designer"]))
        self.assertIn("HTML", skills)
        self.assertEqual(len(history), 2)

    def test_keeps_what_is_already_there(self):
        full = {"first_name": "Janet", "last_name": "Doe-Smith", "home_location": "Bel Air, MD", "state": "MD",
                "primary_job_title": "UX Designer", "job_titles": ["UX Designer"], "skills": ["HTML", "Figma"],
                "work_history": [{"role": "Web Designer", "company": "Example Co", "dates": "2025", "description": "mine"}]}
        added, save = self.fill(full)
        first, last, location, primary, others, skills, history = save.call_args.args
        self.assertEqual((first, last, location, primary), ("Janet", "Doe-Smith", "Bel Air, MD", "UX Designer"))
        self.assertEqual(skills[:2], ["HTML", "Figma"])
        self.assertEqual([h["company"] for h in history], ["Example Co", "Sample Corp"])
        self.assertEqual(history[0]["description"], "mine")  # existing job untouched
        self.assertNotIn("name", added)

    def test_nothing_new_saves_nothing(self):
        found = profile_import.extract(SAMPLE_RESUME_TEXT)
        same = {"first_name": "Jane", "last_name": "Doe", "home_location": "Springfield, MD", "state": "MD",
                "primary_job_title": "Web Designer", "job_titles": ["Web Designer"], "skills": found["skills"],
                "work_history": found["work_history"]}
        added, save = self.fill(same)
        self.assertEqual(added, [])
        save.assert_not_called()


def make_scan_pdf():
    """A PDF page with a drawing but no text, like a scanned certificate."""
    import pymupdf
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.draw_rect(pymupdf.Rect(72, 72, 300, 200), color=(0, 0, 1), fill=(0.9, 0.9, 1))
    data = pdf.tobytes()
    pdf.close()
    return data


class DocumentTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()

    def test_add_each_type(self):
        pdf = documents.add("jane.pdf", make_sample_pdf(), "Old resume")
        docx = documents.add("letter.docx", make_sample_docx())
        txt = documents.add("projects.txt", "Built the Example Co intranet.\n".encode("cp1252"))
        md = documents.add("notes.md", "# Awards\n- Employee of the month".encode("utf-8"))
        self.assertEqual(pdf["label"], "Old resume")
        self.assertEqual(docx["label"], "letter")  # defaults to the file name
        self.assertEqual([d["type"] for d in documents.load()], ["PDF", "DOCX", "TXT", "MD"])
        self.assertIn("Example Co", documents.text(pdf["id"]))
        self.assertIn("EXPERIENCE", documents.text(docx["id"]))
        self.assertIn("intranet", documents.text(txt["id"]))
        self.assertIn("Employee of the month", documents.text(md["id"]))

    def test_rejects_bad_files(self):
        for name, data in (("a.exe", b"MZ"), ("a.pdf", b"not a pdf"), ("a.docx", b"%PDF"), ("a.doc", b"x")):
            with self.assertRaises(documents.DocumentError):
                documents.add(name, data)
        self.assertEqual(documents.load(), [])
        self.assertEqual([p.name for p in (config.DATA_DIR / "documents").glob("*")] if (config.DATA_DIR / "documents").exists() else [], [])

    def test_rename_and_delete(self):
        record = documents.add("a.txt", b"hello")
        documents.rename(record["id"], "Award letter")
        self.assertEqual(documents.load()[0]["label"], "Award letter")
        documents.delete(record["id"])
        self.assertEqual(documents.load(), [])
        self.assertFalse((config.DATA_DIR / "documents" / record["id"]).exists())
        with self.assertRaises(KeyError):
            documents.delete("../current-resume")

    def test_scan_has_no_text_but_has_pictures(self):
        record = documents.add("certificate.pdf", make_scan_pdf(), "Certificate")
        self.assertEqual(record["chars"], 0)
        self.assertEqual(len(documents.page_images(record["id"])), 1)


class DocumentWebAndConnectorTests(ConnectorTests):
    def test_page_upload_view_delete(self):
        import app as web
        client = web.app.test_client()
        with mock.patch.object(jobfinder_db, "get_profile", return_value=PROFILE):
            page = client.get("/").get_data(as_text=True)
            token = re.search(r'name="csrf" value="([^"]+)"', page).group(1)
            from io import BytesIO
            page = client.post("/documents/add", data={"csrf": token, "label": "2024 review",
                                                      "document": (BytesIO(b"Exceeded goals."), "review.txt")},
                               content_type="multipart/form-data", follow_redirects=True).get_data(as_text=True)
        self.assertIn("Added 2024 review.", page)
        doc_id = documents.load()[0]["id"]
        with client.get(f"/documents/{doc_id}/file") as response:
            self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertEqual(client.get("/documents/nope/file").status_code, 404)
        client.post(f"/documents/{doc_id}/delete", data={"csrf": token})
        self.assertEqual(documents.load(), [])

    def test_connector_lists_and_reads(self):
        self.assertIn("No reference documents", self.text(self.call("list_documents")))
        text_doc = documents.add("review.txt", b"Exceeded every goal in 2024.", "2024 review")
        scan = documents.add("certificate.pdf", make_scan_pdf(), "Certificate")
        listed = self.text(self.call("list_documents"))
        self.assertIn("2024 review", listed)
        self.assertIn(scan["id"], listed)
        self.assertIn("Exceeded every goal", self.text(self.call("get_document", {"document_id": text_doc["id"]})))
        result = self.call("get_document", {"document_id": scan["id"]})
        self.assertEqual([b.type for b in result.content].count("image"), 1)
        self.assertIn("No document", self.text(self.call("get_document", {"document_id": "nope"})))

    def test_long_documents_are_cut(self):
        record = documents.add("long.txt", b"x" * (documents.MAX_TEXT_CHARS + 500))
        out = self.text(self.call("get_document", {"document_id": record["id"]}))
        self.assertIn("[Cut at", out)


HISTORY_DOC = """JANE DOE
123 Example St
Springfield, MD 21000
Phone: 555-555-0100
Employment History
Front End Web Developer
05/2015 - 10/2024
Acme Widgets
1 Main St, Springfield, MD
Built landing pages with HTML and CSS.
Sales Associate
11/2013 - 01/2014
Best Buy
2 Mall Rd, Timonium,
MD
Sold computers. Used Photoshop for store signs.
Cashier
09/2008 - 04/2009
Martin's Food Market
3 Town Center
Springfield, MD
Handled cash.
Cashier
04/2009 - 11/2010
Redner's Warehouse Markets
4 Market St
Received money.
"""

CONTACTS_DOC = """Acme Widgets
1 Main St
Springfield
MD
21000
555-555-0199
May 2015 - Oct 2024
Sam Lee
Manager
"""


class DocumentSuggestionTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()
        documents.add("history.txt", HISTORY_DOC.encode("utf-8"), "Employment History")
        documents.add("contacts.txt", CONTACTS_DOC.encode("utf-8"), "Job Addresses")
        self.profile = dict(PROFILE, work_history=[{"role": "Front End Web Developer", "company": "Acme Widgets Inc",
                                                    "dates": "May 2015 - Oct 2024", "description": ""}], skills=["HTML"])

    def test_jobs_and_skills_not_in_the_profile_are_suggested(self):
        import document_import
        found = document_import.suggestions(self.profile)
        self.assertEqual([(j["role"], j["company"]) for j in found["jobs"]],
                         [("Sales Associate", "Best Buy"), ("Cashier", "Martin's Food Market"), ("Cashier", "Redner's Warehouse Markets")])
        self.assertEqual(found["jobs"][0]["description"], "Sold computers. Used Photoshop for store signs.")  # no address left over
        self.assertEqual(found["jobs"][1]["description"], "Handled cash.")
        self.assertEqual(found["jobs"][0]["source"], "Employment History")
        self.assertIn("CSS", found["skills"])
        self.assertNotIn("HTML", found["skills"])  # already in the profile

    def test_only_ticked_items_are_saved(self):
        import app as web
        web.app.config["TESTING"] = True
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            page = client.get("/").get_data(as_text=True)
            self.assertIn("Found in Your Reference Documents", page)
            client.post("/profile/add-from-documents", data={"csrf": "t", "job": ["1"], "skill": ["CSS"]})
        history, skills = save.call_args.args[6], save.call_args.args[5]
        self.assertEqual([w["company"] for w in history], ["Acme Widgets Inc", "Martin's Food Market"])  # newest first
        self.assertEqual(skills, ["HTML", "CSS"])

    def test_nothing_ticked_saves_nothing(self):
        import app as web
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile/add-from-documents", data={"csrf": "t"})
        save.assert_not_called()

ADDRESS_DOC = """Acme Widgets
1 Main St # 2
Springfield
MD
21000
555-555-0100
May 2015 - Oct 2024
Sam Lee
Senior Manager of Design
slee@acmewidgets.com
https://www.example-org.com/
Example Org
9 Side Rd Suite 3
Shelbyville
MD
21001
555-555-0111
April 2014 - April 2015
Pat Kim
CEO
pkim@example-org.com
Sam Lee
555-555-0199
Senior Manager of Design
slee@acmewidgets.com
Robin Diaz
Marketing Technology
555-555-0122
rdiaz@acmewidgets.com
Lee Park
555-555-0133
lee.park@example.net
Education History
"""


class EducationTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()

    def test_the_profile_form_saves_schools(self):
        import app as web
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        form = {"csrf": "t", "first_name": "Jane", "last_name": "Doe", "wh_count": "0", "ed_count": "3",
                "ed0_school": "Example University", "ed0_degree": "Bachelor's Degree", "ed0_major": "Graphic Design",
                "ed0_gpa": "3.5", "ed0_start_date": "2010-09", "ed0_end_date": "2013-05",
                "ed1_school": "Example College", "ed1_degree": "Made-up", "ed1_start_date": "2024-01", "ed1_current": "1",
                "ed2_school": ""}
        with mock.patch.object(jobfinder_db, "get_profile", return_value=PROFILE), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile", data=form)
        education = save.call_args.args[7]
        self.assertEqual(education[0]["gpa"], "3.5")
        self.assertEqual(education[1]["end_date"], "Present")
        self.assertEqual([e["degree"] for e in jobfinder_db.clean_education(education)], ["Bachelor's Degree", "Other"])  # blank school dropped

    def test_zip_and_links_are_saved_like_job_finder(self):
        import app as web
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        form = {"csrf": "t", "first_name": "Jane", "last_name": "Doe", "home_zip": "21009", "linkedin_url": "linkedin.com/in/example",
                "portfolio_url": "javascript:alert(1)"}
        with mock.patch.object(jobfinder_db, "get_profile", return_value=PROFILE), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile", data=form)
        self.assertEqual(jobfinder_db.clean_contact(save.call_args.args[8]),
                         {"home_zip": "21009", "linkedin_url": "https://linkedin.com/in/example", "portfolio_url": ""})

    def test_a_profile_changed_in_job_finder_is_not_overwritten(self):
        import app as web
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        old = jobfinder_db.profile_version(PROFILE)
        newer = dict(PROFILE, skills=PROFILE["skills"] + ["Figma"])  # changed on the Dashboard after this page loaded
        with mock.patch.object(jobfinder_db, "get_profile", return_value=newer), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            page = client.post("/profile", data={"csrf": "t", "profile_version": old, "first_name": "Jane", "last_name": "Doe"},
                               follow_redirects=True).get_data(as_text=True)
        save.assert_not_called()
        self.assertIn("changed in Job Finder after this page loaded", page)
        with mock.patch.object(jobfinder_db, "get_profile", return_value=newer), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile", data={"csrf": "t", "profile_version": jobfinder_db.profile_version(newer),
                                          "first_name": "Jane", "last_name": "Doe"})
        save.assert_called_once()

    def test_other_saves_leave_education_alone(self):
        calls = []

        class Cursor:
            def execute(self, sql, params=()):
                calls.append(sql)

            def fetchone(self):
                return {"table": "x"}

            def fetchall(self):
                return [{"Field": f} for f in ("company", "role", "dates", "description")]

        from contextlib import contextmanager

        @contextmanager
        def fake_cursor(commit=False):
            yield Cursor()

        with mock.patch.object(jobfinder_db, "_cursor", fake_cursor):
            jobfinder_db.save_profile("Jane", "Doe", "", "", [], [], [])
        self.assertFalse(any("user_profile_education" in sql for sql in calls))

    def test_the_page_shows_the_degree_choices(self):
        import app as web
        with mock.patch.object(jobfinder_db, "get_profile", return_value=dict(PROFILE, education=[
                {"school": "Example University", "degree": "Bachelor's Degree", "major": "Design", "minor": "", "start_date": "2010-09",
                 "end_date": "2013-05", "gpa": ""}])):
            page = web.app.test_client().get("/").get_data(as_text=True)
        self.assertIn("Bachelor&#39;s Degree in Design, Example University", page)
        self.assertIn('<option value="Doctorate">Doctorate</option>', page)


class SectionEndTests(unittest.TestCase):
    TEXT = """Employment History
Server
05/2005 - 03/2007
Example Hall
1 Road Rd,
Springfield, MD
Served food to guests.
Education History
05/24/2013
Example University
Bachelor's Degree
References
Kim Lee, Office Clerk
Example Markets
555-555-0144
Supervisor reference known for 3 year(s).
"""

    def test_the_last_job_stops_at_the_next_section(self):
        import document_import
        jobs = document_import.jobs_in(self.TEXT)
        self.assertEqual([(j["role"], j["company"], j["description"]) for j in jobs], [("Server", "Example Hall", "Served food to guests.")])

    def test_a_name_and_title_on_one_line_is_a_person(self):
        import document_import
        people = document_import.contacts_in(self.TEXT)["people"]
        self.assertEqual([(p["name"], p["title"], p["company"], p["phone"]) for p in people],
                         [("Kim Lee", "Office Clerk", "Example Markets", "555-555-0144")])
        self.assertIn("Supervisor reference", people[0]["note"])


class ContactListTests(unittest.TestCase):
    def setUp(self):
        for folder in (config.OUTPUT_DIR, config.DATA_DIR):
            shutil.rmtree(folder, ignore_errors=True)
        config.ensure_dirs()
        documents.add("addresses.txt", ADDRESS_DOC.encode("utf-8"), "Job Addresses")
        self.profile = dict(PROFILE, work_history=[
            {"role": "Front End Web Developer", "company": "Acme Widgets LLC", "dates": "2015", "description": "", "phone": "555-555-0000"},
            {"role": "Junior Web Designer", "company": "Example Org", "dates": "2014", "description": ""}])

    def test_missing_job_details_are_found_and_filled_ones_are_kept(self):
        import document_import
        details = document_import.detail_suggestions(self.profile)
        self.assertEqual(details[0]["details"], {"street": "1 Main St # 2", "city": "Springfield", "state": "MD", "zip": "21000",
                                                 "supervisor_name": "Sam Lee", "supervisor_title": "Senior Manager of Design",
                                                 "supervisor_email": "slee@acmewidgets.com",
                                                 "supervisor_phone": "555-555-0199"})  # his own number; the job's phone was already filled in
        self.assertEqual(details[1]["details"]["website"], "https://www.example-org.com/")

    def test_people_become_possible_references_linked_to_the_users_job(self):
        import document_import
        people = {p["name"]: p for p in document_import.reference_suggestions(self.profile, [{"name": "Lee Park"}])}
        self.assertEqual(sorted(people), ["Lee Park", "Pat Kim", "Robin Diaz", "Sam Lee"])  # no headings
        self.assertTrue(people["Lee Park"]["is_reference"])  # already saved: still listed, marked, so they can also be a supervisor
        placed = dict(self.profile, work_history=[dict(self.profile["work_history"][0], supervisor_name="Lee Park")]
                      + self.profile["work_history"][1:])
        names = [p["name"] for p in document_import.reference_suggestions(placed, [{"name": "Lee Park"}])]
        self.assertNotIn("Lee Park", names)  # a supervisor and a reference: placed both ways, so off the list
        self.assertEqual(people["Sam Lee"]["phone"], "555-555-0199")  # merged from his second listing
        self.assertEqual(people["Sam Lee"]["relationship"], "Supervisor")
        self.assertEqual(people["Robin Diaz"]["user_job"], "Front End Web Developer at Acme Widgets LLC")  # from the email's domain
        self.assertEqual(people["Robin Diaz"]["relationship"], "")

    def test_a_dismissed_person_or_job_detail_is_not_suggested_again(self):
        import app as web
        import document_import
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile):
            page = client.get("/").get_data(as_text=True)
            self.assertIn('value="reference:robin diaz"', page)
            self.assertIn('formaction="/suggestions/dismiss"', page)
            client.post("/suggestions/dismiss", data={"csrf": "t", "dismiss": "reference:robin diaz"})
            client.post("/suggestions/dismiss", data={"csrf": "t", "dismiss": "detail:front end web developer|acme widgets llc"})
            self.assertEqual(client.post("/suggestions/dismiss", data={"csrf": "t", "dismiss": "nonsense:x"}).status_code, 400)
        names = [p["name"] for p in document_import.reference_suggestions(self.profile, [])]
        self.assertNotIn("Robin Diaz", names)
        self.assertIn("Sam Lee", names)
        self.assertEqual([d["company"] for d in document_import.detail_suggestions(self.profile)], ["Example Org"])

    def test_one_click_adds_a_found_person_to_references(self):
        import app as web
        import document_import
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        pat = next(p for p in document_import.reference_suggestions(self.profile, []) if p["name"] == "Pat Kim")
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile):
            page = client.post("/references/add-suggested", data={"csrf": "t", "key": pat["key"]}, follow_redirects=True).get_data(as_text=True)
        self.assertEqual([(r["name"], r["job_title"], r["relationship"]) for r in references.load()], [("Pat Kim", "CEO", "Supervisor")])
        self.assertIn("Added Pat Kim to your References.", page)
        self.assertIn("✓ In your References", page)  # still shown, so they can also be set as a supervisor
        self.assertNotIn('name="key" value="%s">Add to References' % pat["key"], page)

    def test_a_suggested_person_or_saved_reference_becomes_a_jobs_supervisor(self):
        import app as web
        import document_import
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        robin = next(p for p in document_import.reference_suggestions(self.profile, []) if p["name"] == "Robin Diaz")
        self.assertEqual(robin["job_index"], 0)  # linked to the Acme job by her email
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile):
            page = client.get("/").get_data(as_text=True)
        self.assertIn(f'name="sup_job_{robin["key"]}"', page)
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile/set-supervisor", data={"csrf": "t", "person": f"suggestion:{robin['key']}", f"sup_job_{robin['key']}": "1"})
        history = save.call_args.args[6]
        self.assertEqual((history[1]["supervisor_name"], history[1]["supervisor_title"], history[1]["supervisor_email"]),
                         ("Robin Diaz", "Marketing Technology", "rdiaz@acmewidgets.com"))  # the job chosen, not the matched one
        self.assertNotIn("supervisor_name", history[0])
        saved = references.add({"name": "Lee Park", "job_title": "Owner", "phone": "555-555-0133", "user_job": "Junior Web Designer at Example Org"})
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            page = client.get("/").get_data(as_text=True)
            self.assertIn(f'name="sup_job_{saved["id"]}"><option value="">Choose a job…</option><option value="0">', page)
            client.post("/profile/set-supervisor", data={"csrf": "t", "person": f"reference:{saved['id']}", f"sup_job_{saved['id']}": "1"})
        self.assertEqual((save.call_args.args[6][1]["supervisor_name"], save.call_args.args[6][1]["supervisor_title"]), ("Lee Park", "Owner"))
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            client.post("/profile/set-supervisor", data={"csrf": "t", "person": f"reference:{saved['id']}"})  # no job chosen
        save.assert_not_called()

    def test_profile_links_are_shown(self):
        import app as web
        profile = dict(self.profile, linkedin_url="https://www.linkedin.com/in/example", portfolio_url="https://example.com/")
        with mock.patch.object(jobfinder_db, "get_profile", return_value=profile):
            page = web.app.test_client().get("/").get_data(as_text=True)
        self.assertIn('href="https://www.linkedin.com/in/example"', page)
        self.assertIn(">example.com</a>", page)
        self.assertIn('<details class="add-reference-panel"', page)

    def test_ticked_details_and_references_are_saved(self):
        import app as web
        import document_import
        client = web.app.test_client()
        with client.session_transaction() as session:
            session["csrf"] = "t"
        key = next(p["key"] for p in document_import.reference_suggestions(self.profile, []) if p["name"] == "Robin Diaz")
        with mock.patch.object(jobfinder_db, "get_profile", return_value=self.profile), \
                mock.patch.object(jobfinder_db, "save_profile") as save:
            page = client.get("/").get_data(as_text=True)
            self.assertIn("People in Your Documents", page)  # not under References: the user chooses where each one goes
            self.assertLess(page.index('id="found-people"'), page.index('id="references"'))
            self.assertNotIn("Possible References", page)
            client.post("/profile/add-from-documents", data={"csrf": "t", "detail": ["0"], "reference": [key]})
        saved = {w["company"]: w for w in save.call_args.args[6]}
        self.assertEqual(saved["Acme Widgets LLC"]["supervisor_name"], "Sam Lee")
        self.assertEqual(saved["Acme Widgets LLC"]["phone"], "555-555-0000")  # left as it was
        self.assertNotIn("supervisor_name", saved["Example Org"])  # not ticked
        self.assertEqual([(r["name"], r["user_job"]) for r in references.load()],
                         [("Robin Diaz", "Front End Web Developer at Acme Widgets LLC")])


class ForeignHostTests(unittest.TestCase):
    def test_requests_for_another_name_are_refused(self):
        import app as web  # noqa: E402

        client = web.app.test_client()
        self.assertEqual(client.get("/", base_url="http://evil.example:5001").status_code, 400)
        self.assertEqual(client.get("/", base_url="http://localhost.evil.example").status_code, 400)


def tearDownModule():
    shutil.rmtree(TEMP, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
