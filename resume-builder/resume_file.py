"""The user's current resume: saving the upload, reading its text and measuring its layout."""
import json
import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import pymupdf as fitz

import config
import design_report

INFO_FILE = "info.json"
SUBSET_PREFIX = re.compile(r"^[A-Z]{6}\+")
# PostScript names in PDFs ("TimesNewRomanPSMT", "Calibri-Bold") back to family names.
FONT_NAMES = {
    "timesnewroman": "Times New Roman", "times": "Times New Roman", "arial": "Arial",
    "helvetica": "Arial", "calibri": "Calibri", "cambria": "Cambria", "georgia": "Georgia",
    "garamond": "Garamond", "ebgaramond": "Garamond", "verdana": "Verdana", "tahoma": "Tahoma",
    "segoeui": "Segoe UI", "trebuchetms": "Trebuchet MS", "bookantiqua": "Book Antiqua",
    "centurygothic": "Century Gothic", "palatinolinotype": "Palatino Linotype", "palatino": "Palatino Linotype",
    "aptos": "Aptos", "lato": "Lato", "roboto": "Roboto", "opensans": "Open Sans",
}
SERIF_FAMILIES = {"Times New Roman", "Cambria", "Georgia", "Garamond", "Book Antiqua", "Palatino Linotype"}


class ResumeFileError(ValueError):
    pass


def family_name(raw):
    name = SUBSET_PREFIX.sub("", raw or "")
    name = re.split(r"[-,]", name)[0]
    key = re.sub(r"(PSMT|PS|MT|Std|Pro)$", "", name).replace(" ", "").lower()
    return FONT_NAMES.get(key, name)


def _hex(color_int):
    return "#{:06X}".format(color_int & 0xFFFFFF)


def _is_dark(hex_color):
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    return max(r, g, b) < 70 or (abs(r - g) < 20 and abs(g - b) < 20 and max(r, g, b) < 110)


def measure_pdf_layout(pdf_path):
    """Measure fonts, sizes, colours and alignment from the first page of a PDF."""
    with fitz.open(pdf_path) as doc:
        if doc.page_count == 0:
            raise ResumeFileError("The PDF has no pages.")
        page = doc[0]
        width = page.rect.width
        lines = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if spans:
                    lines.append(spans)
        rules = [d for d in page.get_drawings()
                 if d["rect"].height < 2.5 and d["rect"].width > width * 0.5]
    if not lines:
        raise ResumeFileError("No text was found in the PDF. Is it a scanned image?")

    weights = Counter()
    fonts = Counter()
    for spans in lines:
        for s in spans:
            weights[round(s["size"] * 2) / 2] += len(s["text"])
            fonts[family_name(s["font"])] += len(s["text"])
    body_size = weights.most_common(1)[0][0]
    largest = max((s for spans in lines for s in spans), key=lambda s: s["size"])
    name_center = (largest["bbox"][0] + largest["bbox"][2]) / 2
    left = min(s["bbox"][0] for spans in lines for s in spans)

    headings = []
    for spans in lines:
        text = "".join(s["text"] for s in spans).strip()
        first = spans[0]
        bold = bool(first["flags"] & 16) or "bold" in first["font"].lower()
        if spans[0] is largest or len(text) > 40 or not re.search(r"[A-Za-z]", text):
            continue
        if (first["size"] >= body_size + 0.75) or (bold and text.isupper()):
            headings.append((text, first))
    heading_size = Counter(round(h[1]["size"] * 2) / 2 for h in headings).most_common(1)
    accent_colors = Counter(_hex(h[1]["color"]) for h in headings)
    accent_colors[_hex(largest["color"])] += 2
    accent = next((c for c, _ in accent_colors.most_common() if not _is_dark(c)), None)
    text_color = Counter(_hex(s["color"]) for spans in lines for s in spans).most_common(1)[0][0]
    family = fonts.most_common(1)[0][0]
    return {
        "font_family": family,
        "font_kind": "serif" if family in SERIF_FAMILIES else "sans",
        "body_size": body_size,
        "name_size": round(largest["size"] * 2) / 2,
        "heading_size": heading_size[0][0] if heading_size else body_size + 1,
        "accent_color": accent or text_color,
        "text_color": text_color,
        "name_align": "center" if abs(name_center - width / 2) < width * 0.08 else "left",
        "heading_case": "upper" if headings and sum(h[0].isupper() for h in headings) * 2 >= len(headings) else "title",
        "heading_rule": len(rules) >= 2,
        "margin_in": max(0.4, min(1.25, round(left / 72 * 20) / 20)),
    }


def pdf_text(pdf_path):
    with fitz.open(pdf_path) as doc:
        return "\n".join(page.get_text("text") for page in doc).strip()


def page_images(pdf_path, max_pages=2, dpi=110):
    with fitz.open(pdf_path) as doc:
        return [page.get_pixmap(dpi=dpi).tobytes("png") for page in list(doc)[:max_pages]]


def docx_to_pdf(docx_path, pdf_path):
    """Use a private Word instance (never the user's open Word window) to print the file to PDF."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(docx_path), ReadOnly=True, AddToRecentFiles=False)
        try:
            document.SaveAs2(str(pdf_path), FileFormat=17)  # wdFormatPDF
        finally:
            document.Close(False)
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()


def docx_text_and_layout(docx_path):
    """Fallback when Word cannot convert the file: read text and styles with python-docx."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = docx.Document(str(docx_path))
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in paragraphs)
    for table in document.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
    normal = document.styles["Normal"].font
    body_size = normal.size.pt if normal.size else 11.0
    family = normal.name or "Calibri"

    def size(p):
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        return max(sizes) if sizes else (p.style.font.size.pt if p.style.font.size else body_size)

    name_para = paragraphs[0] if paragraphs else None
    headings = [p for p in paragraphs[1:] if len(p.text) <= 40 and (
        p.style.name.lower().startswith("heading") or size(p) > body_size + 0.5
        or (p.runs and all(r.bold for r in p.runs if r.text.strip()) and p.text.isupper()))]
    colors = Counter()
    for p in headings + ([name_para] if name_para else []):
        for r in p.runs:
            if r.font.color is not None and r.font.color.type is not None and r.font.color.rgb is not None:
                colors["#" + str(r.font.color.rgb)] += 1
    accent = next((c for c, _ in colors.most_common() if not _is_dark(c)), "#222222")
    rule = any(p._p.pPr is not None and p._p.pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr") is not None for p in headings)
    family = family_name(family)
    layout = {
        "font_family": family,
        "font_kind": "serif" if family in SERIF_FAMILIES else "sans",
        "body_size": body_size,
        "name_size": size(name_para) if name_para else body_size * 2,
        "heading_size": size(headings[0]) if headings else body_size + 1,
        "accent_color": accent,
        "text_color": "#222222",
        "name_align": "center" if name_para is not None and name_para.alignment == WD_ALIGN_PARAGRAPH.CENTER else "left",
        "heading_case": "upper" if headings and sum(p.text.isupper() for p in headings) * 2 >= len(headings) else "title",
        "heading_rule": rule,
        "margin_in": round(document.sections[0].left_margin.inches, 2) if document.sections else 0.75,
    }
    return text, layout


def _check_signature(data, suffix):
    if suffix == ".pdf" and not data.startswith(b"%PDF"):
        raise ResumeFileError("That file does not look like a PDF.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise ResumeFileError("That file does not look like a Word .docx file.")


def save_upload(filename, data):
    """Replace the current resume with a new upload and measure it. Returns the info dict."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".pdf", ".docx"):
        raise ResumeFileError("Upload a PDF or a Word .docx file. Older .doc files: open in Word and Save As .docx.")
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise ResumeFileError("That file is larger than 10 MB.")
    _check_signature(data, suffix)
    config.ensure_dirs()
    staging = config.UPLOAD_DIR.with_name("current-resume.new")
    shutil.rmtree(staging, ignore_errors=True)
    staging.mkdir(parents=True)
    original = staging / f"resume{suffix}"
    original.write_bytes(data)

    notes = []
    preview_pdf = original if suffix == ".pdf" else staging / "resume-preview.pdf"
    if suffix == ".docx":
        try:
            docx_to_pdf(original, preview_pdf)
        except Exception as error:  # Word missing, busy or blocked: fall back to python-docx
            notes.append(f"Word could not make a preview ({error.__class__.__name__}); layout read from the .docx styles.")
            preview_pdf = None
    try:
        facts = None
        if preview_pdf is not None:
            text, layout = pdf_text(preview_pdf), measure_pdf_layout(preview_pdf)
            facts = _design_facts(preview_pdf)
            layout.update(design_report.layout_from_facts(facts, layout))
            with fitz.open(preview_pdf) as doc:
                pages = doc.page_count
        else:
            (text, layout), pages = docx_text_and_layout(original), None
    except ResumeFileError:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    except Exception as error:
        shutil.rmtree(staging, ignore_errors=True)
        raise ResumeFileError(f"Could not read that file: {error}") from error

    info = {
        "original_name": Path(filename).name[:200],
        "file": original.name,
        "preview_pdf": preview_pdf.name if preview_pdf is not None else "",
        "uploaded_at": datetime.now().isoformat(timespec="seconds"),
        "pages": pages,
        "layout": layout,
        "design": _design_record(facts, layout),
        "notes": notes,
    }
    (staging / "text.txt").write_text(text, encoding="utf-8")
    (staging / INFO_FILE).write_text(json.dumps(info, indent=2), encoding="utf-8")
    shutil.rmtree(config.UPLOAD_DIR, ignore_errors=True)
    staging.rename(config.UPLOAD_DIR)
    return info


def _design_facts(pdf_path):
    try:
        return design_report.analyze_pdf(pdf_path)
    except Exception:  # an odd PDF just gets the plain description from its measured settings
        return None


def _design_record(facts, layout):
    rows = design_report.describe(facts, layout)
    return {"facts": facts, "rows": rows, "notes": design_report.notes_text(rows)}


def reanalyze_design():
    """Look at the uploaded file again (also fills in the design of a résumé uploaded before this existed)."""
    info = current_info()
    if not info:
        return None
    pdf = config.UPLOAD_DIR / (info.get("preview_pdf") or "")
    facts = _design_facts(pdf) if info.get("preview_pdf") and pdf.exists() else None
    info["layout"].update(design_report.layout_from_facts(facts, info["layout"]))
    info["design"] = _design_record(facts, info["layout"])
    (config.UPLOAD_DIR / INFO_FILE).write_text(json.dumps(info, indent=2), encoding="utf-8")
    return info


def current_design():
    """{'facts', 'rows', 'notes'} for the uploaded résumé, looked at once if it has not been yet."""
    info = current_info()
    if not info:
        return None
    if "design" not in info:
        info = reanalyze_design()
    return info["design"]


def current_info():
    path = config.UPLOAD_DIR / INFO_FILE
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def current_text():
    path = config.UPLOAD_DIR / "text.txt"
    return path.read_text(encoding="utf-8") if path.exists() else ""


def current_page_images():
    info = current_info()
    if not info or not info.get("preview_pdf"):
        return []
    return page_images(config.UPLOAD_DIR / info["preview_pdf"])


def current_file_path():
    info = current_info()
    return config.UPLOAD_DIR / info["file"] if info else None
