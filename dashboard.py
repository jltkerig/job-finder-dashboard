from datetime import datetime, timedelta, timezone
import gzip
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import threading
import secrets
import re
import json
import shutil
import tempfile
import time
from urllib.parse import quote, urlparse
from zoneinfo import ZoneInfo

import mysql.connector
from jobfinder import db, paths
from dotenv import load_dotenv
from flask import Flask, abort, jsonify, redirect, render_template, request, send_from_directory, session
from mysql.connector import Error
from jobfinder.profiles.profile_tools import (AMBIGUOUS_SKILLS, SKILL_ALIASES, fit_score, normalize_skills, resume_skill_suggestions,
                           parse_work_history, refresh_listing_skills, resume_suggestions, skill_demand,
                           uploaded_resume)
from jobfinder.profiles.onet_data import occupation_skill_suggestions, proper_title, related_title_suggestions, spelling_fix, title_matches
from jobfinder.profiles.places import city_matches
from jobfinder.profiles.travel import describe as describe_trip
from jobfinder.sources.job_listings import NON_JOB_PATH, is_pdf_url
from jobfinder.db_schema import ensure_unique_source_index
from jobfinder.sources.job_feeds import FEED_NAMES
from jobfinder.sources.job_sites import JOB_SITE_NAMES
from jobfinder.records.board_health import STATUSES as HEALTH_STATUSES, read_health
from jobfinder.records.search_skips import TTL_HOURS, latest_decisions
from jobfinder.records.capture_import import CAPTURE_SOURCES, capture_dirs, move_pending, pending_files
from jobfinder.records.job_retention import CLOSED_KEEP_DAYS, SAVED_STATUSES, tidy_closed_jobs
from jobfinder.web.webfiles import (  # noqa: F401  (also used by callers that import these from here)
    APPLICATION_STATUSES,
    BASE_DIR,
    BLOCKED_COMPANIES_FILE,
    BLOCKED_DOMAINS_FILE,
    BLOCK_METADATA_FILE,
    COMPRESSIBLE,
    EXTENSION_FILE,
    EXTENSION_ID,
    GRACEFUL_STOP_SECONDS,
    JOB_FINDER_PATH,
    LOCAL_HOSTS,
    LOCAL_TIMEZONE,
    RECOMMENDED_DOMAINS_FILE,
    RESUME_FOLDER,
    SEARCH_SKIPS_FILE,
    SKIPS_PER_PAGE,
    UPDATE_MARKER,
    UPDATE_SCRIPT,
    UPDATE_SEARCH_DIRS,
    UPDATE_ZIP_PATTERN,
)
from jobfinder.web import webfiles
from jobfinder.web.core import (  # noqa: F401  (also used by callers that import these from here)
    API_PATH_CODES,
    ERROR_CODES,
    api_error,
    app,
    attach_error_codes,
    get_csrf_token,
    handle_forbidden,
    handle_internal_error,
    how_long,
    inject_csrf_token,
    local_time,
    log_error_code,
    make_responses_lighter,
    protect_local_post_requests,
    refuse_foreign_hosts,
)
from jobfinder.web.schema import (  # noqa: F401  (also used by callers that import these from here)
    _safe_db_identifier,
    _schema_ready,
    ensure_keep_column,
    ensure_profile_tables,
    initialize_database,
    tidy_expired_closed_jobs,
)
from jobfinder.web import schema
from jobfinder.web.search_history import (  # noqa: F401  (also used by callers that import these from here)
    collapse_search_history,
    get_search_history,
    record_search_history,
    split_search_titles,
)
from jobfinder.web.profile_store import (  # noqa: F401  (also used by callers that import these from here)
    RELATED_JOB_TITLES,
    listing_skill_demand,
    profile_skill_suggestions,
    refresh_job_fit,
    related_job_title_suggestions,
)
from jobfinder.web import profile_store
from jobfinder.web.blocklists import (  # noqa: F401  (also used by callers that import these from here)
    add_blocked_company,
    add_domain_to_blocklist,
    block_company,
    block_details,
    manage_blocked_company,
    manage_blocked_domain,
    read_block_metadata,
    remove_domain_from_blocklist,
    save_block_metadata,
)
from jobfinder.web import blocklists
from jobfinder.web.listing_queries import (  # noqa: F401  (also used by callers that import these from here)
    add_drive_times,
    add_job_fit,
    get_companies,
    get_dashboard_counts,
    get_kept_companies,
)
from jobfinder.web import listing_queries
from jobfinder.web.tuning import (  # noqa: F401  (also used by callers that import these from here)
    TUNING_FIELDS,
    TUNING_SWITCHES,
    apply_tuning_form,
    read_tuning_settings,
    save_tuning_settings,
    tuning_page,
)
from jobfinder.web.update_checks import (  # noqa: F401  (also used by callers that import these from here)
    app_version,
    check_update,
    find_latest_update_zip,
    install_update,
)
from jobfinder.web.extension_api import (  # noqa: F401  (also used by callers that import these from here)
    extension_distances,
    extension_file,
    extension_fit_profile,
    extension_updates,
    latest_extension_build,
)


APP_VERSION = "1.1.150"

app.config.update(APP_VERSION=APP_VERSION)


scraper_process = None
scraper_mode = None
scraper_last_error = None
scraper_lock = threading.Lock()
scraper_started_at = None
scraper_stopping = False


def scraper_status():
    global scraper_process, scraper_mode, scraper_last_error, scraper_stopping

    with scraper_lock:
        if scraper_process is None:
            return False, None

        exit_code = scraper_process.poll()
        if exit_code is None:
            return True, scraper_mode

        was_stopping = scraper_stopping
        scraper_stopping = False
        if exit_code != 0 and not was_stopping:
            scraper_last_error = (
                f"Job Finder stopped with exit code {exit_code}. "
                f"Details were saved to {webfiles.SCRAPER_LOG_FILE.name}."
            )

        scraper_process = None
        scraper_mode = None
        return False, None


def scraper_is_running():
    running, _ = scraper_status()
    return running


@app.route("/")
def home():
    now = datetime.now(ZoneInfo("America/New_York"))

    print()
    print("=" * 60)
    print("PAGE REFRESH — " f"{now.strftime('%B %d, %Y at %I:%M:%S %p %Z')}")
    print("=" * 60)

    ensure_keep_column()
    schema.ensure_job_tracking_columns()
    companies = get_companies()

    running, mode = scraper_status()
    profile = profile_store.get_user_profile()
    add_job_fit(companies, profile.get("skills", []))
    recent_search = get_search_history(limit=1)
    latest_search_at = recent_search[0].get("searched_at") if recent_search else None
    skipped = []
    if SEARCH_SKIPS_FILE.exists():
        try:
            for item in reversed(list(latest_decisions(SEARCH_SKIPS_FILE).values())):
                # Older searches recorded PDFs as skips; they are now ignored entirely.
                if item.get("reason") != "Passed" and not is_pdf_url(item.get("url")):
                    skipped.append(item)
                if len(skipped) >= 20:
                    break
        except (OSError, ValueError):
            skipped = []

    return render_template(
        "index.html",
        companies=companies,
        skipped=skipped,
        latest_search_at=latest_search_at,
        scraper_running=running,
        scraper_mode=mode or "",
        profile_cities=profile.get("cities", []),
        profile_job_titles=profile.get("job_titles", []),
        profile_state=profile.get("state", ""),
        profile_work_preferences=profile.get("work_preferences", []),
    )


@app.route("/dashboard")
def user_dashboard():
    ensure_keep_column()
    schema.ensure_job_tracking_columns()
    status_filter = request.args.get("status", "").strip()[:30]
    state_filter = request.args.get("state", "").strip()[:100]
    title_filter = request.args.get("title", "").strip()[:255]
    sort_by = request.args.get("sort", "date_desc").strip()[:30]
    companies = get_kept_companies(status_filter, state_filter, title_filter, sort_by)
    profile = profile_store.get_user_profile()
    add_job_fit(companies, profile.get("skills", []))
    add_drive_times(companies, profile.get("home_zip"), profile.get("state") or "")
    resume_skills = resume_skill_suggestions(RESUME_FOLDER, profile.get("skills", []))
    # With no work history saved yet, the form is filled in from the uploaded résumé (nothing is saved until Save Profile).
    resume_history = [] if profile.get("work_history") else parse_work_history(uploaded_resume(RESUME_FOLDER)[1])
    demanded_skills = [{"skill": skill, "count": count} for skill, count in listing_skill_demand(profile.get("skills", []))]
    return render_template(
        "user-dashboard.html",
        companies=companies,
        profile=profile,
        counts=get_dashboard_counts(),
        search_history=get_search_history(),
        skill_suggestions=profile_skill_suggestions(profile),
        resume_source=resume_skills[0], resume_skills=resume_skills[1],
        demanded_skills=demanded_skills,
        resume_history=resume_history,
        fit_updated=request.args.get("fit_updated", type=int),
        filters={"status": status_filter, "state": state_filter, "title": title_filter, "sort": sort_by},
    )


@app.route("/credibility-scores")
def credibility_scores():
    return render_template("credibility-scores.html")


@app.route("/rejected-listings")
def rejected_listings():
    schema.ensure_job_tracking_columns()
    decisions = latest_decisions(SEARCH_SKIPS_FILE)
    search_skips = []
    for event in reversed(list(decisions.values())):
        if event.get("reason") == "Passed" or is_pdf_url(event.get("url")):
            continue
        checked_at = event.get("checked_at") or ""
        next_check = "Next search"
        if checked_at and event.get("reason") in TTL_HOURS:
            try:
                expires = datetime.fromisoformat(checked_at.replace("Z", "+00:00")) + timedelta(hours=TTL_HOURS[event["reason"]])
                next_check = expires.astimezone(ZoneInfo("America/New_York")).strftime("%b %d, %Y %I:%M %p") if expires > datetime.now(timezone.utc) else "Next search"
            except (ValueError, TypeError):
                pass
        search_skips.append({**event, "next_check": next_check})
    # The skipped pages run to thousands, so only one page of them is sent (and searched on the server).
    skip_query = (request.args.get("skip_q") or "").strip()[:200]
    if skip_query:
        needle = skip_query.casefold()
        search_skips = [item for item in search_skips
                        if needle in " ".join(str(item.get(k) or "") for k in ("title", "reason", "url")).casefold()]
    skip_total = len(search_skips)
    skip_pages = max(1, -(-skip_total // SKIPS_PER_PAGE))
    try:
        skip_page = min(max(int(request.args.get("skip_page", 1)), 1), skip_pages)
    except ValueError:
        skip_page = 1
    search_skips = search_skips[(skip_page - 1) * SKIPS_PER_PAGE:skip_page * SKIPS_PER_PAGE]
    return render_template(
        "rejected-listings.html",
        companies=listing_queries.get_rejected_companies(),
        search_skips=search_skips, skip_total=skip_total, skip_page=skip_page, skip_pages=skip_pages, skip_query=skip_query,
        blocked_domains=blocklists.get_blocked_domains(),
        blocked_companies=blocklists.get_blocked_companies(),
        blocked_domain_info=block_details("domains", blocklists.get_blocked_domains()),
        blocked_company_info=block_details("companies", blocklists.get_blocked_companies()),
        company_notice=request.args.get("company_notice", ""),
        domain_notice=request.args.get("domain_notice", ""),
    )


@app.route("/reject-listing/<int:company_id>", methods=["POST"])
def reject_listing(company_id):
    schema.ensure_job_tracking_columns()
    wants_json = request.headers.get("X-Requested-With") == "fetch" or request.accept_mimetypes.best == "application/json"
    connection = None
    cursor = None
    domain = None
    reason = (request.get_json(silent=True) or {}).get("reason") if request.is_json else request.form.get("reason")
    if reason not in {"wrong_role", "wrong_location", "not_a_job", "duplicate", "other"}:
        reason = "other"
    try:
        connection = db.connect()
        cursor = connection.cursor(dictionary=True)
        cursor.execute("SELECT domain FROM companies WHERE id = %s", (company_id,))
        row = cursor.fetchone()
        if not row:
            if wants_json:
                return api_error("E3201", "The listing could not be found.", 404)
            return redirect(request.referrer or "/")

        domain = row.get("domain")
        cursor.execute(
            """
            UPDATE companies
            SET pre_reject_kept = CASE WHEN is_rejected = 0 THEN is_kept ELSE pre_reject_kept END,
                pre_reject_status = CASE WHEN is_rejected = 0 THEN application_status ELSE pre_reject_status END,
                is_rejected = 1, is_kept = 0, application_status = 'Rejected',
                rejected_at = CURRENT_TIMESTAMP, rejection_reason = %s, rejected_by = 'user'
            WHERE id = %s
            """,
            (reason, company_id),
        )
        connection.commit()
        if wants_json:
            return jsonify({"status": "rejected", "company_id": company_id, "domain": domain})
    except Error as error:
        print("Could not reject listing.")
        print(error)
        if wants_json:
            return api_error("E3202", "Could not reject the listing.", 500)
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()

    return redirect(request.referrer or "/")


@app.route("/restore-rejected/<int:company_id>", methods=["POST"])
def restore_rejected(company_id):
    schema.ensure_job_tracking_columns()
    wants_json = request.headers.get("X-Requested-With") == "fetch" or request.accept_mimetypes.best == "application/json"
    connection = None
    cursor = None
    domain = None
    try:
        connection = db.connect()
        cursor = connection.cursor(dictionary=True)
        cursor.execute("SELECT domain, pre_reject_kept FROM companies WHERE id = %s AND is_rejected = 1", (company_id,))
        row = cursor.fetchone()
        if not row:
            if wants_json:
                return api_error("E3203", "The rejected listing could not be found.", 404)
            return redirect("/rejected-listings")

        domain = row.get("domain")
        # Older rejections have no saved state; they return as unsaved search results.
        cursor.execute(
            """
            UPDATE companies
            SET is_rejected = 0, is_kept = COALESCE(pre_reject_kept, 0),
                application_status = COALESCE(pre_reject_status, 'None'),
                rejected_at = NULL, rejection_reason = NULL, rejected_by = NULL,
                pre_reject_kept = NULL, pre_reject_status = NULL
            WHERE id = %s
            """,
            (company_id,),
        )
        connection.commit()
        if wants_json:
            return jsonify({"status": "restored", "company_id": company_id, "domain": domain,
                            "kept": bool(row.get("pre_reject_kept"))})
    except Error as error:
        print("Could not restore rejected listing.")
        print(error)
        if wants_json:
            return api_error("E3204", "Could not restore the listing.", 500)
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()

    return redirect("/rejected-listings")


@app.route("/block-domain/<int:company_id>", methods=["POST"])
def block_domain(company_id):
    schema.ensure_job_tracking_columns()
    connection = None
    cursor = None
    try:
        connection = db.connect()
        cursor = connection.cursor(dictionary=True)
        cursor.execute("SELECT domain, source_type FROM companies WHERE id = %s", (company_id,))
        row = cursor.fetchone()
        if row and row.get("source_type") in FEED_NAMES:
            return api_error("E3212", "This is the feed domain, not the employer domain. Block the company instead.", 400)
        if not row or not row.get("domain"):
            return api_error("E3210", "No domain was available to block.", 404)
        add_domain_to_blocklist(row["domain"])
        return jsonify({"status": "blocked", "domain": row["domain"]})
    except Error as error:
        return api_error("E3211", "Could not block the domain.", 500)
    finally:
        if cursor is not None: cursor.close()
        if connection is not None and connection.is_connected(): connection.close()


@app.route("/save-profile", methods=["POST"])
def save_profile():
    first_name = request.form.get("first_name", "").strip()[:100]
    last_name = request.form.get("last_name", "").strip()[:100]
    state = request.form.get("state", "").strip()[:100]
    raw_job_titles = request.form.get("job_titles", "")[:5000]
    job_titles = []

    for part in re.split(r"[,\n]+", raw_job_titles):
        title = proper_title(part.strip()[:255])
        if title and title.lower() not in {item.lower() for item in job_titles}:
            job_titles.append(title)

    try:
        cities = json.loads(request.form.get("cities_json", "[]") or "[]")
    except json.JSONDecodeError:
        cities = []
    def read_list(key):
        try:
            value = json.loads(request.form.get(key, "[]"))
            return value if isinstance(value, list) else []
        except (TypeError, ValueError):
            return []
    home_location = request.form.get("home_location", "").strip()[:150]
    home_zip = re.sub(r"\D", "", request.form.get("home_zip", ""))[:5]
    if home_zip and len(home_zip) != 5:
        home_zip = ""  # a half-typed ZIP is dropped rather than saved
    # The primary title has its own box; it is always searched too, so it leads the title list.
    primary = proper_title(request.form.get("primary_job_title", "").strip()[:255])
    if primary:
        job_titles = [primary] + [title for title in job_titles if title.casefold() != primary.casefold()]
    previous_skills = {str(skill).casefold() for skill in profile_store.get_user_profile().get("skills", [])}
    avatar = request.form.get("avatar_data", "")
    if avatar and (not re.fullmatch(r"data:image/jpeg;base64,[A-Za-z0-9+/=]+", avatar) or len(avatar) > 550000):
        abort(400)
    profile_store.save_user_profile(first_name, last_name, state, job_titles, cities,
                      home_location=home_location, home_zip=home_zip, primary_job_title=primary,
                      skills=read_list("skills_json"), work_history=read_list("work_history_json"),
                      avatar_data=avatar if "avatar_data" in request.form else None,
                      work_preferences=[value for value in request.form.getlist("work_preferences")
                                        if value in {"Part-time", "Full-time", "Contract", "Freelance / Gig", "Remote", "Hybrid", "Onsite"}])
    # Changed skills: the percentages on the pages follow at once (they are worked out on each load); listings saved
    # earlier are re-read so they also list skills the current skills list recognizes.
    if {str(skill).casefold() for skill in read_list("skills_json")} != previous_skills:
        return redirect(f"/dashboard?fit_updated={refresh_job_fit()}")
    return redirect("/dashboard")


@app.route("/profile/save-title", methods=["POST"])
def save_profile_title():
    data = request.get_json(silent=True) or {}
    title = proper_title(str(data.get("title", "")).strip()[:255])
    if not title:
        return api_error("E3301", "Enter a job title.")
    profile = profile_store.get_user_profile()
    if title.casefold() in {s.casefold() for s in profile["job_titles"]}:
        return jsonify({"status": "already_saved"})
    profile["job_titles"].append(title)
    if not profile_store.save_user_profile(profile["first_name"], profile["last_name"], profile["state"],
                             profile["job_titles"], profile["cities"]):
        return api_error("E3302", "Could not save the job title.", 500)
    return jsonify({"status": "saved"})


@app.route("/profile/parse-resume", methods=["POST"])
def parse_resume():
    upload = request.files.get("resume")
    if not upload or not upload.filename:
        return api_error("E3310", "Choose a PDF or Word document.")
    extension = Path(upload.filename).suffix.lower()
    if extension not in (".pdf", ".docx"):
        return api_error("E3311", "Use a PDF or DOCX résumé.")
    data = upload.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        return api_error("E3312", "The résumé must be under 5 MB.")
    try:
        from io import BytesIO
        if extension == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in PdfReader(BytesIO(data)).pages[:20])
        else:
            from docx import Document
            document = Document(BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    except Exception as error:
        print(f"Could not parse résumé: {error}")
        return api_error("E3313", "Could not read that résumé. Try another PDF or DOCX.")
    if not text.strip():
        return api_error("E3314", "No selectable text was found. A scanned image résumé needs OCR.")
    return jsonify({"status": "ok", "suggestions": resume_suggestions(text[:250000])})


@app.route("/save-kept", methods=["POST"])
def save_kept():
    ensure_keep_column()

    data = request.get_json(silent=True) or {}
    company_ids = data.get("company_ids", [])

    if not isinstance(company_ids, list):
        return jsonify({"status": "error", "message": "Invalid company list."}), 400

    cleaned_ids = []
    for company_id in company_ids:
        try:
            cleaned_ids.append(int(company_id))
        except (TypeError, ValueError):
            continue

    if not cleaned_ids:
        return jsonify({"status": "error", "message": "Select at least one result to keep."}), 400

    connection = None
    cursor = None

    try:
        connection = db.connect()
        cursor = connection.cursor()
        placeholders = ",".join(["%s"] * len(cleaned_ids))
        cursor.execute(
            f"UPDATE companies SET is_kept = 1, application_status = 'Saved' WHERE id IN ({placeholders})",
            cleaned_ids,
        )
        connection.commit()
        return jsonify({"status": "saved", "count": cursor.rowcount})
    except Error as error:
        return jsonify({"status": "error", "message": str(error)}), 500
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()


@app.route("/unsave-kept/<int:company_id>", methods=["POST"])
def unsave_kept(company_id):
    schema.ensure_job_tracking_columns()
    connection = None
    cursor = None
    try:
        connection = db.connect()
        cursor = connection.cursor()
        cursor.execute(
            "UPDATE companies SET is_kept = 0, application_status = 'None' WHERE id = %s",
            (company_id,),
        )
        connection.commit()
        return jsonify({"status": "unsaved", "company_id": company_id})
    except Error as error:
        return api_error("E3102", "Could not unsave the listing.", 500)
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()


def _finish_graceful_stop(process):
    """Wait for a search to clean up after a stop request; force it only as a fallback."""
    try:
        process.wait(timeout=GRACEFUL_STOP_SECONDS)
    except subprocess.TimeoutExpired:
        log_error_code("E2106", "Job Finder did not stop in time; forcing it and stopping SearXNG.")
        process.kill()
        process.wait(timeout=10)
        # The killed process could not run its own cleanup.
        try:
            subprocess.run(["docker", "stop", "searxng"], capture_output=True, timeout=60,
                           creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        except (OSError, subprocess.TimeoutExpired) as error:
            log_error_code("E2106", f"Could not stop SearXNG: {error}")
    finally:
        webfiles.STOP_REQUEST_FILE.unlink(missing_ok=True)


@app.route("/stop-search", methods=["POST"])
def stop_search():
    global scraper_process, scraper_mode, scraper_last_error, scraper_stopping
    with scraper_lock:
        if scraper_process is None or scraper_process.poll() is not None:
            scraper_process = None
            scraper_mode = None
            return jsonify({"status": "stopped", "message": "No Job Finder process was running."})
        if scraper_mode in ("search", "replacement"):
            # Searches own SearXNG and Docker, so let them shut those down themselves.
            if not scraper_stopping:
                scraper_stopping = True
                try:
                    webfiles.STOP_REQUEST_FILE.touch()
                except OSError as error:
                    scraper_stopping = False
                    log_error_code("E2105", f"Could not request a stop: {error}")
                    return api_error("E2105", "Could not stop the current Job Finder action.", 500)
                threading.Thread(target=_finish_graceful_stop, args=(scraper_process,), daemon=True).start()
            return jsonify({"status": "stopping"})
        try:
            scraper_process.terminate()
            try:
                scraper_process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                scraper_process.kill()
                scraper_process.wait(timeout=5)
            scraper_process = None
            scraper_mode = None
            scraper_last_error = None
            return jsonify({"status": "stopped"})
        except OSError as error:
            return api_error("E2105", "Could not stop the current Job Finder action.", 500)


@app.route("/update-kept/<int:company_id>", methods=["POST"])
def update_kept(company_id):
    schema.ensure_job_tracking_columns()
    application_status = request.form.get("application_status", "None").strip()
    notes = request.form.get("notes", "").strip()[:5000]
    if application_status not in APPLICATION_STATUSES:
        application_status = "None"

    connection = None
    cursor = None
    try:
        connection = db.connect()
        cursor = connection.cursor()
        cursor.execute(
            "UPDATE companies SET application_status = %s, notes = %s WHERE id = %s AND is_kept = 1",
            (application_status, notes, company_id),
        )
        connection.commit()
    except Error as error:
        print("Could not update saved result.")
        print(error)
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()
    # The #job anchor reopens this card on the Dashboard, where saved jobs start collapsed.
    return redirect((request.referrer or "/dashboard") + f"#job-{company_id}")


@app.route("/delete-kept/<int:company_id>", methods=["POST"])
def delete_kept(company_id):
    ensure_keep_column()
    connection = None
    cursor = None

    try:
        connection = db.connect()
        cursor = connection.cursor()
        cursor.execute(
            "DELETE FROM companies WHERE id = %s AND is_kept = 1",
            (company_id,),
        )
        connection.commit()
    except Error as error:
        print()
        print("Could not delete saved result.")
        print(error)
    finally:
        if cursor is not None:
            cursor.close()
        if connection is not None and connection.is_connected():
            connection.close()

    return redirect("/dashboard")


def validate_search_criteria(job_title, state, cities=None):
    job_title = (job_title or "").strip()[:1000]
    state = (state or "").strip()[:100]

    if not job_title:
        return None, None, None, "Enter a job title."

    if not state:
        return None, None, None, "Enter a state."

    cleaned_cities = []
    allowed_radii = {5, 10, 15, 20, 30, 50}
    if isinstance(cities, list):
        seen = set()
        for item in cities[:25]:
            if not isinstance(item, dict):
                continue
            city = str(item.get("city", "")).strip()[:150]
            try:
                radius = int(item.get("radius", 50))
            except (TypeError, ValueError):
                radius = 50
            if city and radius in allowed_radii and city.lower() not in seen:
                seen.add(city.lower())
                cleaned_cities.append({"city": city, "radius": radius})

    return job_title, state, cleaned_cities, None


def launch_search_process(job_title, state, cities=None, mode="search"):
    global scraper_process, scraper_mode, scraper_started_at

    job_title, state, cities, validation_error = validate_search_criteria(job_title, state, cities)
    if validation_error:
        return jsonify({"status": "error", "message": validation_error}), 400

    with scraper_lock:
        if scraper_process is not None and scraper_process.poll() is None:
            return jsonify({"status": "already_running", "mode": scraper_mode}), 200

        if not JOB_FINDER_PATH.exists():
            return jsonify({"status": "error", "message": "job_finder.py was not found."}), 500

        # Keep Search and User Profile on the same saved job-title/location data,
        # but only once this search is actually going to start.
        if mode != "replacement":
            try:
                profile = profile_store.get_user_profile()
                titles = []
                for part in job_title.split(","):
                    title = part.strip()[:255]
                    if title and title.lower() not in {item.lower() for item in titles}:
                        titles.append(title)
                profile_store.save_user_profile(
                    profile.get("first_name", ""),
                    profile.get("last_name", ""),
                    state,
                    titles,
                    cities,
                )
            except Exception as error:
                log_error_code("E3301", f"Could not sync search criteria to profile: {error}")

        try:
            creationflags = 0
            if os.name == "nt":
                creationflags = subprocess.CREATE_NO_WINDOW

            global scraper_last_error
            scraper_last_error = None
            # A leftover request from an earlier stop would end this search immediately.
            webfiles.STOP_REQUEST_FILE.unlink(missing_ok=True)
            with webfiles.SCRAPER_LOG_FILE.open("a", encoding="utf-8") as log_file:
                log_file.write(
                    f"\n=== {datetime.now().isoformat(timespec='seconds')} | {mode} | "
                    f"{job_title} | {state} ===\n"
                )
                log_file.flush()
                scraper_process = subprocess.Popen(
                    [
                        sys.executable,
                        "-u",
                        str(JOB_FINDER_PATH),
                        "--job-title",
                        job_title,
                        "--state",
                        state,
                        "--cities-json",
                        json.dumps(cities or []),
                    ] + (["--max-new", "1"] if mode == "replacement" else []),
                    cwd=str(BASE_DIR),
                    creationflags=creationflags,
                    stdout=log_file,
                    stderr=subprocess.STDOUT,
                )
            scraper_mode = mode
            scraper_started_at = time.monotonic()
            if mode != "replacement":
                record_search_history(job_title, state, cities)
        except OSError as error:
            scraper_process = None
            scraper_mode = None
            return jsonify({"status": "error", "message": str(error)}), 500

    return jsonify({"status": "started"}), 202


@app.route("/job-title-matches")
def job_title_matches():
    """Type-ahead for job title boxes: real job titles that match what has been typed so far."""
    return jsonify({"matches": title_matches((request.args.get("q") or "")[:100])})


_home_state_cache = {"value": "", "at": 0.0}


def _home_state():
    """The profile's home state ("MD"), re-read at most once a minute so typing doesn't hit the database."""
    if time.monotonic() - _home_state_cache["at"] > 60:
        _home_state_cache.update(value=(profile_store.get_user_profile().get("state") or "")[:2], at=time.monotonic())
    return _home_state_cache["value"]


@app.route("/city-matches")
def city_matches_route():
    """Type-ahead for the city boxes: U.S. places that start with what was typed, near home first."""
    return jsonify({"matches": city_matches((request.args.get("q") or "")[:100], _home_state())})


@app.route("/job-title-suggestions")
def job_title_suggestions():
    titles = (request.args.get("titles") or "").strip()[:1000]
    return jsonify({"suggestions": related_job_title_suggestions(titles)})

@app.route("/start-search", methods=["POST"])
def start_search():
    data = request.get_json(silent=True) or {}
    return launch_search_process(
        data.get("job_title"),
        data.get("state"),
        data.get("cities"),
        mode="search",
    )


@app.route("/replace-result", methods=["POST"])
def replace_result():
    history = get_search_history(limit=1)
    if not history:
        return jsonify({"status": "no_results", "message": "Start a search to save criteria before requesting a replacement."}), 200
    search = history[0]
    try:
        cities = json.loads(search.get("cities_json") or "[]")
    except (TypeError, ValueError):
        cities = []
    return launch_search_process(search.get("job_title"), search.get("state"), cities, mode="replacement")


@app.route("/refresh-search", methods=["POST"])
def refresh_search():
    """Recheck the result rows currently shown without changing search criteria."""
    global scraper_process, scraper_mode, scraper_started_at, scraper_last_error
    data = request.get_json(silent=True) or {}
    ids = data.get("company_ids")
    if not isinstance(ids, list) or len(ids) > 500 or any(type(value) is not int or value <= 0 for value in ids):
        return jsonify({"status": "error", "message": "Invalid result selection."}), 400
    ids = list(dict.fromkeys(ids))
    if not ids:
        return jsonify({"status": "no_results", "message": "There are no displayed listings to refresh."})
    with scraper_lock:
        if scraper_process is not None and scraper_process.poll() is None:
            return jsonify({"status": "already_running", "mode": scraper_mode}), 200
        if not JOB_FINDER_PATH.exists():
            return jsonify({"status": "error", "message": "job_finder.py was not found."}), 500
        try:
            flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            scraper_last_error = None
            with webfiles.SCRAPER_LOG_FILE.open("a", encoding="utf-8") as log_file:
                log_file.write(f"\n=== {datetime.now().isoformat(timespec='seconds')} | update-existing ===\n")
                log_file.flush()
                scraper_process = subprocess.Popen(
                    [sys.executable, "-u", str(JOB_FINDER_PATH), "--update-existing", "--update-ids", ",".join(map(str, ids))],
                    cwd=str(BASE_DIR), creationflags=flags, stdout=log_file, stderr=subprocess.STDOUT,
                )
            scraper_mode = "refresh"
            scraper_started_at = time.monotonic()
        except OSError as error:
            scraper_process = None
            scraper_mode = None
            return jsonify({"status": "error", "message": str(error)}), 500
    return jsonify({"status": "started"}), 202


@app.route("/update-existing", methods=["POST"])
def update_existing():
    global scraper_process, scraper_mode, scraper_started_at

    with scraper_lock:
        if scraper_process is not None and scraper_process.poll() is None:
            return jsonify({"status": "already_running", "mode": scraper_mode}), 200

        if not JOB_FINDER_PATH.exists():
            return jsonify({"status": "error", "message": "job_finder.py was not found."}), 500

        try:
            creationflags = 0
            if os.name == "nt":
                creationflags = subprocess.CREATE_NO_WINDOW

            global scraper_last_error
            scraper_last_error = None
            with webfiles.SCRAPER_LOG_FILE.open("a", encoding="utf-8") as log_file:
                log_file.write(
                    f"\n=== {datetime.now().isoformat(timespec='seconds')} | update-existing ===\n"
                )
                log_file.flush()
                scraper_process = subprocess.Popen(
                    [sys.executable, "-u", str(JOB_FINDER_PATH), "--update-existing"],
                    cwd=str(BASE_DIR),
                    creationflags=creationflags,
                    stdout=log_file,
                    stderr=subprocess.STDOUT,
                )
            scraper_mode = "update"
            scraper_started_at = time.monotonic()
        except OSError as error:
            scraper_process = None
            scraper_mode = None
            return jsonify({"status": "error", "message": str(error)}), 500

    return jsonify({"status": "started"}), 202


@app.route("/captures/pending")
def captures_pending():
    """Capture files the Web Job Scraper extension left in Downloads, for the import prompt."""
    downloads_dir, searches_dir = capture_dirs(read_tuning_settings(), BASE_DIR)
    try:
        files = pending_files(downloads_dir)
    except OSError as error:
        return api_error("E2120", f"Could not read {downloads_dir}: {error}", 500)
    return jsonify({
        "count": len(files),
        "jobs": sum(item["jobs"] for item in files),
        "from": str(downloads_dir),
        "to": str(searches_dir),
        # Changes whenever a file is added or rewritten, so "Not now" only hides the prompt until something new arrives.
        "signature": ";".join(f"{item['relative']}@{int(item['modified'])}" for item in files),
    })


@app.route("/captures/import", methods=["POST"])
def captures_import():
    """Move the waiting capture files into web-job-scraper\\searches, then run job_finder.py --import-captures."""
    global scraper_process, scraper_mode, scraper_started_at, scraper_last_error
    downloads_dir, searches_dir = capture_dirs(read_tuning_settings(), BASE_DIR)
    with scraper_lock:
        if scraper_process is not None and scraper_process.poll() is None:
            return jsonify({"status": "already_running", "mode": scraper_mode}), 200
        try:
            moved = move_pending(downloads_dir, searches_dir)
        except OSError as error:
            return api_error("E2121", f"Could not move the capture files to {searches_dir}: {error}", 500)
        try:
            flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
            scraper_last_error = None
            with webfiles.SCRAPER_LOG_FILE.open("a", encoding="utf-8") as log_file:
                log_file.write(f"\n=== {datetime.now().isoformat(timespec='seconds')} | import | "
                               f"{len(moved)} capture file(s) ===\n")
                log_file.flush()
                scraper_process = subprocess.Popen(
                    [sys.executable, "-u", str(JOB_FINDER_PATH), "--import-captures"],
                    cwd=str(BASE_DIR), creationflags=flags, stdout=log_file, stderr=subprocess.STDOUT,
                )
            scraper_mode = "import"
            scraper_started_at = time.monotonic()
        except OSError as error:
            scraper_process = None
            scraper_mode = None
            return api_error("E2121", f"Could not start the import: {error}", 500)
    return jsonify({"status": "started", "moved": len(moved)}), 202


def import_progress_from_log():
    if not webfiles.SCRAPER_LOG_FILE.exists():
        return "Preparing the import"
    try:
        with webfiles.SCRAPER_LOG_FILE.open("rb") as log_file:
            size = log_file.seek(0, 2)
            log_file.seek(max(0, size - 65536))
            tail = log_file.read().decode("utf-8", errors="replace")
    except OSError:
        return "Preparing the import"
    current = tail.rsplit("| import |", 1)[-1]
    matches = list(re.finditer(r"Importing (\d+)/(\d+): ([^\r\n]+)", current))
    if matches:
        latest = matches[-1]
        return f"Checking {latest.group(1)} of {latest.group(2)}: {latest.group(3)[:70]}"
    return "Preparing the import"


def update_progress_from_log():
    if not webfiles.SCRAPER_LOG_FILE.exists():
        return "Preparing existing results"
    try:
        with webfiles.SCRAPER_LOG_FILE.open("rb") as log_file:
            size = log_file.seek(0, 2)
            log_file.seek(max(0, size - 65536))
            tail = log_file.read().decode("utf-8", errors="replace")
    except OSError:
        return "Preparing existing results"
    current = tail.rsplit("| update-existing ===", 1)[-1]
    matches = list(re.finditer(r"Updating (\d+)/(\d+): ([^\r\n]+)", current))
    if matches:
        latest = matches[-1]
        completed = "Existing row updated in place." in current[latest.end():] or "Could not update this existing row." in current[latest.end():]
        verb = "Checked" if completed else "Checking"
        return f"{verb} {latest.group(1)} of {latest.group(2)}: {latest.group(3)[:70]}"
    found = re.findall(r"Found (\d+) existing results to verify", current)
    if found:
        return f"Found {found[-1]} existing results to check"
    return "Preparing existing results"


def search_progress_from_log(mode):
    fallback = "Starting replacement search" if mode == "replacement" else "Starting search"
    if not webfiles.SCRAPER_LOG_FILE.exists():
        return {"progress": fallback, "passed": 0, "checked": 0, "saved": 0, "limit": None}
    try:
        with webfiles.SCRAPER_LOG_FILE.open("rb") as log_file:
            size = log_file.seek(0, 2)
            log_file.seek(max(0, size - 2097152))
            tail = log_file.read().decode("utf-8", errors="replace")
    except OSError:
        return {"progress": fallback, "passed": 0, "checked": 0, "saved": 0, "limit": None}
    markers = list(re.finditer(r"\| " + re.escape(mode) + r" \|[^\r\n]*===", tail))
    if not markers:
        return {"progress": fallback, "passed": 0, "checked": 0, "saved": 0, "limit": None}
    current = tail[markers[-1].end():]
    candidates = list(re.finditer(r"Checking result (\d+): ([^\r\n]+)", current))
    passed_matches = list(re.finditer(r"Passed validation: (\d+)", current))
    passed = int(passed_matches[-1].group(1)) if passed_matches else 0
    # "Passed" counts every job that passed, including ones already in the list; only new saves count toward the limit.
    saved_matches = list(re.finditer(r"Saved (?:job |viable company \()(\d+)/(\d+)", current))
    saved = int(saved_matches[-1].group(1)) if saved_matches else 0
    limit = int(saved_matches[-1].group(2)) if saved_matches else None
    events = list(re.finditer(r"^(Checking result \d+: |Skipped \([^\r\n]+?\): |Passed validation: \d+ · |Saved lead for review: )([^\r\n]+)", current, re.M))
    if events:
        latest = events[-1]
        progress = (latest.group(1) + latest.group(2))[:220]
    else:
        progress = "Searching for results"
    checked = int(candidates[-1].group(1)) if candidates else 0
    return {"progress": progress, "passed": passed, "checked": checked, "saved": saved, "limit": limit}


@app.route("/search-status")
def search_status():
    running, mode = scraper_status()
    elapsed = int(time.monotonic() - scraper_started_at) if running and scraper_started_at else 0
    activity = search_progress_from_log(mode) if running and mode in ("search", "replacement") else None
    stop_reason = None
    if not running and webfiles.SCRAPER_LOG_FILE.exists():
        try:
            with webfiles.SCRAPER_LOG_FILE.open("rb") as log_file:
                size = log_file.seek(0, 2)
                log_file.seek(max(0, size - 8192))
                reasons = re.findall(r"Stop reason: ([^\r\n]+)", log_file.read().decode("utf-8", errors="replace"))
                stop_reason = reasons[-1] if reasons else None
        except OSError:
            pass
    return jsonify({
        "running": running,
        "mode": mode,
        "stopping": running and scraper_stopping,
        "error": scraper_last_error,
        "progress": (update_progress_from_log() if mode in ("update", "refresh") else import_progress_from_log() if mode == "import"
                     else activity["progress"] if activity else None) if running else None,
        "passed": activity["passed"] if activity else None,
        "saved": activity["saved"] if activity else None,
        "limit": activity["limit"] if activity else None,
        "checked": activity["checked"] if activity else None,
        "elapsed_seconds": elapsed,
        "stop_reason": stop_reason,
    })


if __name__ == "__main__":
    initialize_database()
    ensure_keep_column()
    schema.ensure_job_tracking_columns()
    ensure_profile_tables()
    tidy_expired_closed_jobs()
    app.run(host="127.0.0.1", port=5000, debug=False, use_reloader=False)
