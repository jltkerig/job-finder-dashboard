"""Saving your profile, reading a résumé for it, and the job-title and city suggestions as you type."""

import json
from pathlib import Path
import re

from flask import abort, jsonify, redirect, request

from jobfinder.profiles.onet_data import proper_title, title_matches
from jobfinder.profiles.places import city_matches
from jobfinder.profiles.profile_tools import resume_suggestions
from jobfinder.web import profile_store
from jobfinder.web.core import api_error, app
from jobfinder.web.profile_store import refresh_job_fit, related_job_title_suggestions, related_skills_for
from jobfinder.web.search_control import _home_state


@app.route("/save-profile", methods=["POST"])
def save_profile():
    first_name = request.form.get("first_name", "").strip()[:100]
    last_name = request.form.get("last_name", "").strip()[:100]
    state = request.form.get("state", "").strip()[:100]
    raw_job_titles = request.form.get("job_titles", "")[:5000]
    job_titles = []

    for part in re.split(r"[,\n]+", raw_job_titles):
        title = proper_title(part.strip()[:255])
        if title and title.lower() not in {item.lower() for item in job_titles}:
            job_titles.append(title)

    try:
        cities = json.loads(request.form.get("cities_json", "[]") or "[]")
    except json.JSONDecodeError:
        cities = []
    def read_list(key):
        try:
            value = json.loads(request.form.get(key, "[]"))
            return value if isinstance(value, list) else []
        except (TypeError, ValueError):
            return []
    home_location = request.form.get("home_location", "").strip()[:150]
    home_zip = re.sub(r"\D", "", request.form.get("home_zip", ""))[:5]
    if home_zip and len(home_zip) != 5:
        home_zip = ""  # a half-typed ZIP is dropped rather than saved
    # The primary title has its own box; it is always searched too, so it leads the title list.
    primary = proper_title(request.form.get("primary_job_title", "").strip()[:255])
    if primary:
        job_titles = [primary] + [title for title in job_titles if title.casefold() != primary.casefold()]
    saved_profile = profile_store.get_user_profile()
    sent_version = request.form.get("profile_version", "")
    if sent_version and sent_version != profile_store.profile_version(saved_profile):
        return redirect("/dashboard?profile_changed=1")  # changed elsewhere since this page loaded: don't overwrite it
    previous_skills = {str(skill).casefold() for skill in saved_profile.get("skills", [])}
    avatar = request.form.get("avatar_data", "")
    if avatar and (not re.fullmatch(r"data:image/jpeg;base64,[A-Za-z0-9+/=]+", avatar) or len(avatar) > 550000):
        abort(400)
    profile_store.save_user_profile(first_name, last_name, state, job_titles, cities,
                      home_location=home_location, home_zip=home_zip, primary_job_title=primary,
                      skills=read_list("skills_json"), work_history=read_list("work_history_json"),
                      education=read_list("education_json"),
                      linkedin_url=request.form.get("linkedin_url", ""), portfolio_url=request.form.get("portfolio_url", ""),
                      avatar_data=avatar if "avatar_data" in request.form else None,
                      work_preferences=[value for value in request.form.getlist("work_preferences")
                                        if value in {"Part-time", "Full-time", "Contract", "Freelance / Gig", "Remote", "Hybrid", "Onsite"}])
    if "avoid_terms" in request.form:
        profile_store.save_avoid_terms(re.split(r"[,\n]+", request.form.get("avoid_terms", "")[:5000]))
    # Changed skills: the percentages on the pages follow at once (they are worked out on each load); listings saved
    # earlier are re-read so they also list skills the current skills list recognizes.
    if {str(skill).casefold() for skill in read_list("skills_json")} != previous_skills:
        return redirect(f"/dashboard?fit_updated={refresh_job_fit()}")
    return redirect("/dashboard")


@app.route("/profile/save-title", methods=["POST"])
def save_profile_title():
    data = request.get_json(silent=True) or {}
    title = proper_title(str(data.get("title", "")).strip()[:255])
    if not title:
        return api_error("E3301", "Enter a job title.")
    profile = profile_store.get_user_profile()
    if title.casefold() in {s.casefold() for s in profile["job_titles"]}:
        return jsonify({"status": "already_saved"})
    profile["job_titles"].append(title)
    if not profile_store.save_user_profile(profile["first_name"], profile["last_name"], profile["state"],
                             profile["job_titles"], profile["cities"]):
        return api_error("E3302", "Could not save the job title.", 500)
    return jsonify({"status": "saved", "profile_version": profile_store.profile_version(profile_store.get_user_profile())})


@app.route("/profile/parse-resume", methods=["POST"])
def parse_resume():
    upload = request.files.get("resume")
    if not upload or not upload.filename:
        return api_error("E3310", "Choose a PDF or Word document.")
    extension = Path(upload.filename).suffix.lower()
    if extension not in (".pdf", ".docx"):
        return api_error("E3311", "Use a PDF or DOCX résumé.")
    data = upload.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        return api_error("E3312", "The résumé must be under 5 MB.")
    try:
        from io import BytesIO
        if extension == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text(extraction_mode="layout") or "" for page in PdfReader(BytesIO(data)).pages[:20])
        else:
            from docx import Document
            document = Document(BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [cell.text for table in document.tables for row in table.rows for cell in row.cells])
    except Exception as error:
        print(f"Could not parse résumé: {error}")
        return api_error("E3313", "Could not read that résumé. Try another PDF or DOCX.")
    if not text.strip():
        return api_error("E3314", "No selectable text was found. A scanned image résumé needs OCR.")
    return jsonify({"status": "ok", "suggestions": resume_suggestions(text[:250000])})


@app.route("/job-title-matches")
def job_title_matches():
    """Type-ahead for job title boxes: real job titles that match what has been typed so far."""
    return jsonify({"matches": title_matches((request.args.get("q") or "")[:100])})


@app.route("/city-matches")
def city_matches_route():
    """Type-ahead for the city boxes: U.S. places that start with what was typed, near home first."""
    return jsonify({"matches": city_matches((request.args.get("q") or "")[:100], _home_state())})


@app.route("/job-title-suggestions")
def job_title_suggestions():
    titles = (request.args.get("titles") or "").strip()[:1000]
    return jsonify({"suggestions": related_job_title_suggestions(titles)})


@app.route("/skill-related")
def skill_related():
    """Skills that go with the one being typed on the Dashboard ("HTML" -> CSS, Responsive Design ...)."""
    return jsonify({"skills": related_skills_for((request.args.get("q") or "").strip()[:80])})
